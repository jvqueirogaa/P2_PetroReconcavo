# -*- coding: utf-8 -*-
# Matriz de cobertura das normas (EC + NBC TA 200 a 600) e registro de versoes das normas
# Preparador: matricula 211059370 | 24/10/2026 | Revisao: ______ (nao preenchido)
# Entradas: papeis_trabalho/PT_*.md e a lista de itens de requisito de cada norma:
#   - se existir _work/normas/REQ_xxx.txt (requisitos copiados do consolidado da disciplina), extrai os numeros dos itens;
#   - senao, usa matriz_cobertura/ITENS_REQUISITOS_EXTRAIDOS.csv (gerado na primeira execucao e versionado no repositorio).
# Saida: matriz_cobertura/MATRIZ_COBERTURA_NORMAS.xlsx (abas: Matriz_por_norma, Itens_requisitos, Nao_se_aplica, Versoes_normas)
# Uso: python codigo/gerar_matriz_cobertura.py
import re, os, glob, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

NORMAS = ['EC', '200', '210', '220', '230', '240', '250', '260', '265', '300', '315', '320', '330', '402', '450',
          '500', '501', '505', '510', '520', '530', '540', '550', '560', '570', '580', '600']
TITULO = {'EC': 'Estrutura Conceitual para Trabalhos de Asseguracao', '200': 'Objetivos gerais do auditor', '210': 'Termos do trabalho',
          '220': 'Gestao da qualidade da auditoria', '230': 'Documentacao', '240': 'Fraude', '250': 'Leis e regulamentos',
          '260': 'Comunicacao com a governanca', '265': 'Deficiencias de controle interno', '300': 'Planejamento',
          '315': 'Identificacao e avaliacao de riscos', '320': 'Materialidade', '330': 'Respostas aos riscos',
          '402': 'Organizacao prestadora de servicos', '450': 'Avaliacao das distorcoes', '500': 'Evidencia de auditoria',
          '501': 'Itens especificos (estoque, litigios, segmentos)', '505': 'Confirmacoes externas', '510': 'Saldos iniciais',
          '520': 'Procedimentos analiticos', '530': 'Amostragem', '540': 'Estimativas contabeis', '550': 'Partes relacionadas',
          '560': 'Eventos subsequentes', '570': 'Continuidade operacional', '580': 'Representacoes formais', '600': 'Auditoria de grupo'}
# Versao/data conforme o texto do consolidado da disciplina (_work/normas/NBC_TA_xxx.txt); "2009*" = versao original sem data no texto
VERSAO = {'EC': ('NBC TA Estrutura Conceitual', '20/11/2015'), '200': ('NBC TA 200 (R1)', '19/08/2016'), '210': ('NBC TA 210 (R1)', '19/08/2016'),
          '220': ('NBC TA 220 (R3)', '18/11/2021'), '230': ('NBC TA 230 (R1)', '19/08/2016'), '240': ('NBC TA 240 (R1) + Revisao NBC 3', '19/08/2016'),
          '250': ('NBC TA 250', '07/02/2019'), '260': ('NBC TA 260 (R2)', '17/06/2016'), '265': ('NBC TA 265', '2009*'),
          '300': ('NBC TA 300 (R1)', '19/08/2016'), '315': ('NBC TA 315 (R2)', '19/08/2021'), '320': ('NBC TA 320 (R1)', '19/08/2016'),
          '330': ('NBC TA 330 (R1) + Revisao NBC 3', '19/08/2016'), '402': ('NBC TA 402', '2009*'), '450': ('NBC TA 450 (R1) + Revisao NBC 3', '19/08/2016'),
          '500': ('NBC TA 500 (R1) + Revisao NBC 3', '19/08/2016'), '501': ('NBC TA 501', '27/11/2009'), '505': ('NBC TA 505', '2009*'),
          '510': ('NBC TA 510 (R1)', '19/08/2016'), '520': ('NBC TA 520', '27/11/2009'), '530': ('NBC TA 530', '2009*'),
          '540': ('NBC TA 540 (R2)', '17/10/2019'), '550': ('NBC TA 550', '2009*'), '560': ('NBC TA 560 (R1)', '19/08/2016'),
          '570': ('NBC TA 570', '17/06/2016'), '580': ('NBC TA 580 (R1)', '19/08/2016'), '600': ('NBC TA 600 (R2)', '25/10/2023')}
VIGENCIA_NOTA = {'600': 'Periodos iniciados em ou apos 01/01/2024 -> aplicavel a 2025', '250': 'Periodos iniciados em ou apos 31/12/2018',
                 '240': 'Dupla numeracao no consolidado (Revisao NBC 3): PT cita a numeracao nova e a antiga entre parenteses'}

def pt_file(n):
    f = glob.glob(f'papeis_trabalho/PT_{n}_*.md'); return f[0] if f else None

def itens_req(n):
    p = f'_work/normas/REQ_{n}.txt'
    if not os.path.exists(p): return None
    txt = open(p, encoding='utf-8', errors='ignore').read().replace('\x00', '')
    if n == '250':   # REQ_250 ficou incompleto na extracao: usa o texto integral, secao de requisitos (itens 13 a 30)
        t2 = open('_work/normas/NBC_TA_250.txt', encoding='utf-8', errors='ignore').read().replace('\x00', '')
        out = []
        for m in re.finditer(r'(?m)^\s?(\d{1,2})\.\s+(\S.{0,110})', t2):
            k = int(m.group(1))
            if 13 <= k <= 30 and k not in [x for x, _ in out]: out.append((k, m.group(2).strip()))
        return sorted(out)
    out, vistos = [], set()
    for m in re.finditer(r'(?m)^\s?(\d{1,2})\.\s+(\S.{0,110})', txt):
        k = int(m.group(1))
        if k in vistos: continue
        vistos.add(k); out.append((k, m.group(2).strip()))
    return sorted(out)

CSV = 'matriz_cobertura/ITENS_REQUISITOS_EXTRAIDOS.csv'
os.makedirs('matriz_cobertura', exist_ok=True)
itens = {}
if all(os.path.exists(f'_work/normas/REQ_{n}.txt') for n in NORMAS if n != 'EC'):
    for n in NORMAS:
        if n == 'EC': continue
        itens[n] = [(k, t) for k, t in itens_req(n) if not (n == '550' and k < 10)]   # REQ_550 inclui a introducao (itens 1-9)
    with open(CSV, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh, delimiter=';'); w.writerow(['norma', 'item', 'inicio_do_texto'])
        for n, l in itens.items():
            for k, t in l: w.writerow([n, k, t])
else:
    for row in csv.DictReader(open(CSV, encoding='utf-8'), delimiter=';'):
        itens.setdefault(row['norma'], []).append((int(row['item']), row['inicio_do_texto']))

def numeros_citados(txt):
    """Numeros de itens citados: cabecalho 'Norma e itens' (com faixas 'x a y', 'x-y', 'x–y') + mencoes 'item/itens ...'."""
    cit = set()
    def faixa(s):
        s = re.sub(r'\b[Aa]\d+\b', ' ', s)                 # ignora itens de aplicacao (A1...)
        for a, b in re.findall(r'(\d{1,2})\s*(?:a|–|-|até)\s*(\d{1,2})\b', s):
            a, b = int(a), int(b)
            if a < b <= 99: cit.update(range(a, b + 1))
        cit.update(int(x) for x in re.findall(r'\b(\d{1,2})\b', s))
    h = re.search(r'\*\*Norma e itens\*\*\s*\|([^\n]*)', txt)
    if h: faixa(h.group(1))
    for m in re.finditer(r'(?i)\bite(?:m|ns)\s+((?:\d{1,2}(?:\s*\([^)]{1,12}\))*(?:\s*(?:,|e|a|–|-|até)\s*)?)+)', txt):
        faixa(m.group(1))
    return cit

H = PatternFill('solid', fgColor='1F4E78'); thin = Side(style='thin', color='BFBFBF'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
OK = PatternFill('solid', fgColor='E2EFDA'); NO = PatternFill('solid', fgColor='FCE4D6')
TOPO = 'P2 Auditoria 1 - PetroReconcavo S.A. (RECV3) - exercicio 2025 | Preparador: matricula 211059370 | 24/10/2026 | Revisao: ______ (nao preenchido)'
def cab(ws, r, cols, larg):
    for j, (c, w) in enumerate(zip(cols, larg), 1):
        x = ws.cell(row=r, column=j, value=c); x.font = Font(bold=True, color='FFFFFF'); x.fill = H; x.border = BD; x.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(j)].width = w
def put(ws, r, vals, fill=None):
    for j, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=j, value=v); x.border = BD; x.alignment = Alignment(wrap_text=True, vertical='top')
        if fill: x.fill = fill

wb = Workbook()
ws = wb.active; ws.title = 'Matriz_por_norma'
ws['A1'] = 'Matriz de cobertura das normas - NBC TA Estrutura Conceitual e NBC TA 200 a 600'; ws['A1'].font = Font(bold=True, size=13); ws['A2'] = TOPO
cab(ws, 4, ['Norma', 'Titulo', 'Versao consultada', 'PT', 'Arquivo', 'Itens citados no cabecalho do PT', 'Qtd. itens de requisito', 'Itens citados', 'Cobertura %',
            '[REALIZADO]', '[SIMULADO]', '[PLANEJADO]', 'Nao se aplica (analisado)', 'Classificacao declarada no cabecalho do PT'], [8, 34, 30, 8, 40, 60, 10, 10, 10, 10, 10, 10, 12, 70])
wi = wb.create_sheet('Itens_requisitos'); wi['A1'] = 'Itens de requisito de cada norma x citacao no PT correspondente'; wi['A1'].font = Font(bold=True, size=13); wi['A2'] = TOPO
cab(wi, 4, ['Norma', 'Item', 'Inicio do texto do requisito', 'PT', 'Citado no PT?'], [8, 6, 90, 8, 12])
wn = wb.create_sheet('Nao_se_aplica'); wn['A1'] = '"Nao se aplica" - cada ocorrencia com a analise que a fundamenta (enunciado: "nao se aplica" exige analise)'; wn['A1'].font = Font(bold=True, size=13); wn['A2'] = TOPO
cab(wn, 4, ['PT', 'Trecho (analise)'], [10, 150])
r, ri, rn = 5, 5, 5
for n in NORMAS:
    f = pt_file(n); txt = open(f, encoding='utf-8').read() if f else ''
    head = re.search(r'\*\*Norma e itens\*\*\s*\|([^\n]*)', txt); head = head.group(1).strip(' |') if head else ''
    cit = numeros_citados(txt)
    lst = itens.get(n) or []
    nc = sum(1 for k, _ in lst if k in cit)
    cobert = (nc / len(lst)) if lst else None
    cl = re.search(r'\*\*Classifica[^*]*\*\*\s*\|([^\n]*)', txt); cl = re.sub(r'[*]', '', cl.group(1)).strip(' |') if cl else (re.search(r'\*\*Natureza\*\*\s*\|([^\n]*)', txt).group(1).strip(' |') if re.search(r'\*\*Natureza\*\*', txt) else '')
    nsa = [l.strip() for l in txt.splitlines() if re.search(r'(?i)n[ãa]o se aplica', l) and 'enunciado' not in l.lower()]
    put(ws, r, [n, TITULO[n], VERSAO[n][0] + ' - ' + VERSAO[n][1], f'PT {n}', f or 'PENDENTE', head[:600], len(lst) if lst else 'n/a (EC)',
                nc if lst else 'n/a', (round(cobert, 3) if cobert is not None else 'n/a'), txt.count('[REALIZADO'), txt.count('[SIMULADO'),
                txt.count('[PLANEJADO'), len(nsa), cl[:500]], OK if (cobert is None or cobert >= 0.999) else NO)
    if cobert is not None: ws.cell(row=r, column=9).number_format = '0%'
    r += 1
    for k, t in lst:
        put(wi, ri, [n, k, t, f'PT {n}', 'SIM' if k in cit else 'NAO'], OK if k in cit else NO); ri += 1
    for l in nsa:
        put(wn, rn, [f'PT {n}', re.sub(r'\s+', ' ', l)[:900]]); rn += 1
ws.cell(row=r + 1, column=1, value='Cobertura = itens de requisito (numerados) da norma citados no PT (cabecalho ou texto). Itens nao citados aparecem em laranja na aba Itens_requisitos e devem ser conferidos.')
ws.cell(row=r + 2, column=1, value='Outros PTs: PT 999 (conclusao integrada, sem opiniao). Classificacao dos procedimentos conforme o enunciado: REALIZADO (evidencia publica), PLANEJADO (falta de acesso), SIMULADO (fins didaticos).')

wv = wb.create_sheet('Versoes_normas'); wv['A1'] = 'Registro das versoes das normas consultadas'; wv['A1'].font = Font(bold=True, size=13); wv['A2'] = TOPO
cab(wv, 4, ['Norma', 'Sigla/versao', 'Data (conforme consolidado)', 'Observacao de vigencia', 'Fonte'], [8, 36, 22, 70, 60])
for i, n in enumerate(NORMAS, 5):
    put(wv, i, [n, VERSAO[n][0], VERSAO[n][1], VIGENCIA_NOTA.get(n, 'Vigente para auditorias de 2025'),
                'Consolidado das normas disponibilizado na disciplina (pasta do projeto); texto em _work/normas/NBC_TA_' + n + '.txt'])
wv.cell(row=6 + len(NORMAS), column=1, value='2009* = versao original publicada em 2009; o texto do consolidado nao traz a data de aprovacao.')
wb.save('matriz_cobertura/MATRIZ_COBERTURA_NORMAS.xlsx')
tot = sum(len(v) for v in itens.values())
print('ok matriz; itens', tot)
for n in NORMAS:
    if n == 'EC': continue
    f = pt_file(n); txt = open(f, encoding='utf-8').read() if f else ''
    cit = numeros_citados(txt); faltam = [k for k, _ in itens.get(n, []) if k not in cit]
    if faltam: print(n, 'nao citados:', faltam)
