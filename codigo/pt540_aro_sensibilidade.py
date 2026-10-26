# -*- coding: utf-8 -*-
# PT 540 - Provisao para abandono de pocos (ARO): recalculos, revisao retrospectiva e sensibilidade
# NBC TA 540 (R2) itens 14, 18(c), 23(c), 24, 26, 28, 29 e 39; NBC TA 320 A11 (sensibilidade); dica da aula (sensibilidade de 1%)
# Preparador: matricula 211059370 | 24/10/2026 | Revisao: ______ (nao preenchido)
#
# O QUE E REALIZADO x SIMULADO
#  - REALIZADO (evidencia publica): movimentacao da provisao (NE 13.1 DFP 2025; NE 14 DF 2024), recalculo da
#    atualizacao (unwinding) pela taxa do ano anterior, revisao retrospectiva da parcela circulante, premissas (NE 3.3).
#  - SIMULADO (fins didaticos): a companhia NAO divulga o cronograma de desembolsos. O script cria um cronograma
#    sintetico de custos (em moeda de 31/12/2025) CALIBRADO para que o valor presente com as premissas da NE 3.3
#    (taxa nominal 10,56% a.a.; inflacao 3,56% a.a.) reproduza exatamente o saldo publicado (R$ 145.192 mil), com a
#    parcela de 2026 igual ao passivo circulante publicado (R$ 4.728 mil). Sobre esse cronograma se calcula a
#    sensibilidade de +/- 1 p.p. Os numeros de sensibilidade sao ILUSTRATIVOS, nao da companhia.
#
# Uso: python codigo/pt540_aro_sensibilidade.py   (gera planilhas/PT540_ARO_SENSIBILIDADE_RECV3.xlsx com FORMULAS)
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# ---------------- 1. Premissas e dados publicos (R$ mil) ----------------
R, PI = 0.1056, 0.0356          # NE 3.3 (2025)
R24, PI24 = 0.1092, 0.0366      # NE 3.3 (2024)
SALDO25, CIRC25 = 145192, 4728  # NE 13.1 consolidado
SALDO24, CIRC24 = 136972, 342   # NE 13.1 / DF 2024 NE 14
MOV25 = dict(constituicao=8383, reversao=-265, atualizacao=14958, baixas=-14856)
ME, MG = 20475, 31500           # PT 320
ANO0, PICO, ULT = 2025, 2040, 2050   # cronograma sintetico: 2026 = parcela circulante; 2027-2050 triangular com pico em 2040
# pico ~ 2025 + reservas 1P / producao anual = 143,3 MMboe / (26,5 mil boe/d x 365) = ~14,8 anos (F14; RA)

def pv(custos, r, pi, choque=1.0, desloc=0):
    """custos: dict ano->custo em moeda de 31/12/2025; fluxo no fim de cada ano; desloc<0 antecipa."""
    s = 0.0
    for ano, c in custos.items():
        t = max(1, ano - ANO0 + desloc)
        s += c * choque * (1 + pi) ** t / (1 + r) ** t
    return s

# ---------------- 2. Calibracao do cronograma sintetico ----------------
anos = list(range(2027, ULT + 1))
peso = {a: (a - 2026) / (PICO - 2026) if a <= PICO else (ULT + 1 - a) / (ULT + 1 - PICO) for a in anos}
c2026 = CIRC25 * (1 + R) / (1 + PI)                        # PV(2026) = parcela circulante
base_nc = sum(peso[a] * (1 + PI) ** (a - ANO0) / (1 + R) ** (a - ANO0) for a in anos)
k = (SALDO25 - CIRC25) / base_nc
custos = {2026: c2026, **{a: k * peso[a] for a in anos}}
assert abs(pv(custos, R, PI) - SALDO25) < 1e-6
nao_desc = sum(custos.values())

# ---------------- 3. Sensibilidade (valores calculados em Python p/ o PT; a planilha refaz por formula) ----------------
cen = [('Base (NE 3.3: 10,56% / 3,56%)', R, PI, 1.0, 0),
       ('Taxa de desconto +1 p.p. (11,56%)', R + .01, PI, 1.0, 0),
       ('Taxa de desconto -1 p.p. (9,56%)', R - .01, PI, 1.0, 0),
       ('Inflacao +1 p.p. (4,56%)', R, PI + .01, 1.0, 0),
       ('Inflacao -1 p.p. (2,56%)', R, PI - .01, 1.0, 0),
       ('Pior combinacao: taxa -1 p.p. e inflacao +1 p.p.', R - .01, PI + .01, 1.0, 0),
       ('Custos +1% (sensibilidade de 1% - aula)', R, PI, 1.01, 0),
       ('Custos +10%', R, PI, 1.10, 0),
       ('Antecipacao de 2 anos no abandono', R, PI, 1.0, -2),
       ('Premissas de 2024 (10,92% / 3,66%) sobre o mesmo cronograma', R24, PI24, 1.0, 0)]
res = []
for n, r, p, ch, d in cen:
    v = pv(custos, r, p, ch, d); dv = v - SALDO25
    res.append(dict(cenario=n, taxa=r, inflacao=p, choque=ch, desloc=d, valor=round(v, 0), delta=round(dv, 0),
                    pct_ME=round(dv / ME, 4), pct_MG=round(dv / MG, 4)))

# ---------------- 4. Planilha com formulas ----------------
wb = Workbook()
H = PatternFill('solid', fgColor='1F4E78'); thin = Side(style='thin', color='BFBFBF'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);-'; PCT = '0.00%'
def cab(ws, r, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = Font(bold=True, color='FFFFFF'); x.fill = H; x.border = BD; x.alignment = Alignment(wrap_text=True)
def lin(ws, r, vals, nfs=None):
    for j, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=j, value=v); x.border = BD
        if nfs and j in nfs: x.number_format = nfs[j]
TOPO = 'PetroReconcavo S.A. - consolidado - R$ mil | Preparador: matricula 211059370 | 24/10/2026 | Revisao: ______ (nao preenchido)'

# Aba 1 - Premissas
ws = wb.active; ws.title = 'Premissas'
ws['A1'] = 'PT 540 - Provisao para abandono de pocos (ARO) - premissas e dados publicos'; ws['A1'].font = Font(bold=True, size=13); ws['A2'] = TOPO
for j, w in enumerate([60, 16, 70], 1): ws.column_dimensions[get_column_letter(j)].width = w
cab(ws, 4, ['Item', 'Valor', 'Fonte / classificacao'])
prem = [('Taxa de desconto nominal 2025', R, 'NE 3.3 F01 - REALIZADO', PCT), ('Inflacao 2025', PI, 'NE 3.3 F01 - REALIZADO', PCT),
        ('Taxa de desconto nominal 2024', R24, 'NE 3.3 F01 - REALIZADO', PCT), ('Inflacao 2024', PI24, 'NE 3.3 F01 - REALIZADO', PCT),
        ('Saldo da provisao 31/12/2025', SALDO25, 'NE 13.1 F01 - REALIZADO', NUM), ('  parcela circulante 31/12/2025', CIRC25, 'NE 13.1 F01', NUM),
        ('Saldo da provisao 31/12/2024', SALDO24, 'NE 13.1 F01 / NE 14 F07', NUM), ('  parcela circulante 31/12/2024', CIRC24, 'NE 14 F07 (DF 2024)', NUM),
        ('Materialidade para execucao (ME)', ME, 'PT 320', NUM), ('Materialidade global (MG)', MG, 'PT 320', NUM)]
for i, (a, b, c, f) in enumerate(prem, 5): lin(ws, i, [a, b, c], {2: f})
# B5 R, B6 PI, B7 R24, B8 PI24, B9 saldo25, B10 circ25, B11 saldo24, B12 circ24, B13 ME, B14 MG

# Aba 2 - Movimentacao e revisao retrospectiva (REALIZADO)
w2 = wb.create_sheet('Retrospectiva_item14')
w2['A1'] = 'Movimentacao 2025 e revisao retrospectiva (NBC TA 540 item 14) - REALIZADO COM EVIDENCIA PUBLICA'; w2['A1'].font = Font(bold=True, size=13); w2['A2'] = TOPO
for j, w in enumerate([60, 16, 16, 60], 1): w2.column_dimensions[get_column_letter(j)].width = w
cab(w2, 4, ['Movimentacao 2025 (NE 13.1 consolidado)', 'Valor', '', 'Observacao'])
mv = [('Saldo 31/12/2024', '=Premissas!B11'), ('Constituicao', MOV25['constituicao']), ('Reversao', MOV25['reversao']),
      ('Atualizacao (unwinding)', MOV25['atualizacao']), ('Baixas (abandonos realizados)', MOV25['baixas'])]
for i, (a, b) in enumerate(mv, 5): lin(w2, i, [a, b], {2: NUM})
lin(w2, 10, ['Saldo recalculado 31/12/2025', '=SUM(B5:B9)', '', 'Deve igualar a NE 13.1'], {2: NUM})
lin(w2, 11, ['Saldo publicado 31/12/2025', '=Premissas!B9', '', ''], {2: NUM})
lin(w2, 12, ['Diferenca', '=B10-B11', '', 'Zero = movimentacao fecha'], {2: NUM})
cab(w2, 14, ['Recalculos', 'Esperado', 'Publicado', 'Leitura'])
lin(w2, 15, ['Atualizacao 2025 = saldo 2024 x taxa nominal 2024', '=Premissas!B11*Premissas!B7', '=B8', 'Unwinding coerente com a taxa de 2024'], {2: NUM, 3: NUM})
lin(w2, 16, ['Baixas 2025 x parcela circulante estimada em 31/12/2024', '=Premissas!B12', '=-B9', 'Baixas realizadas muito acima do previsto como circulante'], {2: NUM, 3: NUM})
lin(w2, 17, ['Multiplo baixas realizadas / circulante previsto', '=C16/B16', '', 'Indicador de falha na estimativa de EPOCA (A55-A60); ver PT 540'], {2: '0.0"x"'})
lin(w2, 18, ['Baixas 2025 x parcela circulante de 31/12/2025', '=Premissas!B10', '=-B9', 'Se o padrao se repetir, o circulante de 2025 pode estar subestimado (classificacao)'], {2: NUM, 3: NUM})
lin(w2, 19, ['Taxa real implicita 2025 = (1+r)/(1+pi)-1', '=(1+Premissas!B5)/(1+Premissas!B6)-1', '', 'Queda da taxa real aumenta a provisao'], {2: PCT})
lin(w2, 20, ['Taxa real implicita 2024', '=(1+Premissas!B7)/(1+Premissas!B8)-1', '', 'Direcao da mudanca de premissa: AUMENTA o passivo (nao indica vies para reduzir)'], {2: PCT})

# Aba 3 - Cronograma sintetico (SIMULADO)
w3 = wb.create_sheet('Cronograma_SIMULADO')
w3['A1'] = 'Cronograma SINTETICO de desembolsos calibrado ao saldo publicado - SIMULADO PARA FINS DIDATICOS'; w3['A1'].font = Font(bold=True, size=13); w3['A2'] = TOPO
w3['A3'] = 'A companhia nao divulga o cronograma. Perfil: 2026 = parcela circulante; 2027-2050 triangular com pico em 2040 (~14,8 anos de reservas 1P / producao). Custo em moeda de 31/12/2025.'
for j, w in enumerate([10, 8, 18, 18, 18, 18, 18, 18, 18, 18, 18, 18], 1): w3.column_dimensions[get_column_letter(j)].width = w
heads = ['Ano', 't', 'Custo (moeda 2025)']
cab(w3, 5, heads + [c['cenario'] for c in res])
# linha 4: parametros de cada cenario (taxa, inflacao, choque, deslocamento) nas colunas D..M
for j, c in enumerate(res, 4):
    col = get_column_letter(j)
    w3[f'{col}4'] = f"r={c['taxa']:.4f}; pi={c['inflacao']:.4f}; x{c['choque']}; desloc {c['desloc']}"
r0 = 6; anos_todos = [2026] + anos
for i, a in enumerate(anos_todos):
    rr = r0 + i
    w3.cell(row=rr, column=1, value=a); w3.cell(row=rr, column=2, value=f'=A{rr}-{ANO0}')
    x = w3.cell(row=rr, column=3, value=round(custos[a], 6)); x.number_format = NUM
    for j, c in enumerate(res, 4):
        t = f'MAX(1,B{rr}+({c["desloc"]}))'
        f = f'=C{rr}*{c["choque"]}*(1+{c["inflacao"]})^{t}/(1+{c["taxa"]})^{t}'
        y = w3.cell(row=rr, column=j, value=f); y.number_format = NUM
rl = r0 + len(anos_todos)
w3.cell(row=rl, column=1, value='TOTAL'); w3.cell(row=rl, column=3, value=f'=SUM(C{r0}:C{rl-1})').number_format = NUM
for j in range(4, 4 + len(res)):
    col = get_column_letter(j); w3.cell(row=rl, column=j, value=f'=SUM({col}{r0}:{col}{rl-1})').number_format = NUM

# Aba 4 - Sensibilidade (resumo, referencia as somas)
w4 = wb.create_sheet('Sensibilidade')
w4['A1'] = 'Sensibilidade da provisao (NBC TA 540 itens 26 e 29; NBC TA 320 A11) - SIMULADO sobre cronograma calibrado'; w4['A1'].font = Font(bold=True, size=13); w4['A2'] = TOPO
for j, w in enumerate([58, 16, 16, 12, 12, 40], 1): w4.column_dimensions[get_column_letter(j)].width = w
cab(w4, 4, ['Cenario', 'Provisao (PV)', 'Variacao vs base', '% da ME', '% da MG', 'Leitura'])
for i, c in enumerate(res, 5):
    col = get_column_letter(4 + i - 5)
    lin(w4, i, [c['cenario'], f"=Cronograma_SIMULADO!{col}{rl}", f'=B{i}-$B$5', f'=C{i}/Premissas!$B$13', f'=C{i}/Premissas!$B$14',
                'acima da ME' if abs(c['delta']) > ME else ('acima do trivial (1.575)' if abs(c['delta']) > 1575 else 'trivial')],
        {2: NUM, 3: NUM, 4: PCT, 5: PCT})
w4.cell(row=6 + len(res), column=1, value='Total nao descontado do cronograma (moeda 2025): ' + f'{nao_desc:,.0f}'.replace(',', '.'))
w4.cell(row=7 + len(res), column=1, value='Valores ILUSTRATIVOS: dependem do cronograma sintetico; a direcao e a ordem de grandeza servem para planejar o teste do item 24 e avaliar a divulgacao (item 26).')

wb.save('planilhas/PT540_ARO_SENSIBILIDADE_RECV3.xlsx')
json.dump(dict(custo_2026=c2026, k=k, nao_descontado=nao_desc, cenarios=res,
               unwinding_esperado=SALDO24 * R24, unwinding_publicado=MOV25['atualizacao'],
               baixas_sobre_circulante_2024=-MOV25['baixas'] / CIRC24, taxa_real_25=(1 + R) / (1 + PI) - 1, taxa_real_24=(1 + R24) / (1 + PI24) - 1),
          open('planilhas/PT540_resultados.json', 'w'), indent=1, ensure_ascii=False)
for c in res: print(f"{c['cenario']:<62} {c['valor']:>10,.0f} {c['delta']:>9,.0f} {c['pct_ME']:>7.1%}")
print('nao descontado', round(nao_desc), 'unwinding', round(SALDO24 * R24), 'baixas/circ24', round(-MOV25['baixas'] / CIRC24, 1))
