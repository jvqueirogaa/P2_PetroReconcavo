# -*- coding: utf-8 -*-
# =============================================================================
# P2 - AUDITORIA 1 (UnB 2026.2) - PetroReconcavo S.A. (RECV3) - exercicio 2025
# Preparador: matricula 211059370
#
# >>> SECAO DE SCRIPTS DE AUDITORIA (testes sobre a BASE SIMULADA) <<<
# Procedimentos classificados como "SIMULADO PARA FINS DIDATICOS": executados
# sobre DADOS SINTETICOS. Nenhum resultado aqui e constatacao sobre a companhia real.
#
# Ordem:  (1) conciliacao base x NE 7 / CVM
#         (2) testes na populacao inteira (T01-T10) - feitos "as cegas" (sem ler o log de erros)
#         (3) amostragem MUS (Gramling, cap. 8; NBC TA 530) - selecao e avaliacao (LSI)
#         (4) SO ENTAO le o log de erros inseridos e compara inseridos x detectados
# Uso:  python codigo/testes_auditoria_base.py     (na raiz P2_PetroReconcavo; apos gerar a base)
# =============================================================================
import os, json
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_BASE = pd.Timestamp('2025-12-31')
B = 'base_simulada'; RES = 'base_simulada/resultados'; os.makedirs(RES, exist_ok=True); os.makedirs('planilhas', exist_ok=True)
SEED_MUS = 211059370 + 530          # semente do ponto de partida aleatorio do MUS (reprodutivel)

# ---------------------------------------------------------------- parametros do auditor - definidos no PT 320 (NBC TA 320 (R1))
# MG  (item 10; A4-A8): referencial = RECEITA LIQUIDA consolidada 2025 (R$ 3.157.609 mil, F01 DRE p.52), porque o lucro antes
#      do imposto e VOLATIL (2024: R$ 363.405 mil; 2025: R$ 776.194 mil, com R$ 366.796 mil de swaps) - item A5 ("quando o lucro
#      antes do imposto ... e volatil, outros referenciais podem ser mais apropriados, como lucro bruto ou receita total").
#      Percentual: 1% (exemplo do item A8 para receita total) -> R$ 31.576 mil, arredondado para baixo: R$ 31.500 mil.
# ME  (item 11; A13): 65% da MG - julgamento (a norma nao fixa percentual): primeira auditoria (sem historico de distorcoes,
#      A13), inconsistencias de divulgacao ja identificadas e risco significativo de impairment -> expectativa de distorcoes.
# IT  da populacao do imobilizado = ME (nao foi fixada materialidade especifica para o imobilizado - item 10/A11).
# RA, RI, RC, IE: modelo de risco do material de MUS (Gramling) e PT 315/330.
PARAM = dict(
    MG=31_500_000.00,          # materialidade para as DCs como um todo (R$) - PT 320
    pct_ME=0.65,               # materialidade para execucao = 65% da MG (PT 320, A13)
    RA=0.05,                   # risco de auditoria aceitavel
    RI=1.00,                   # risco inerente do imobilizado: maximo (estimativas UoP, impairment, capitalizacao - PT 315)
    RC=1.00,                   # risco de controle: maximo (sem teste de controles - sem acesso a entidade)
    pct_IE=0.10,               # incorrecao esperada = 10% da incorrecao toleravel (sem historico - primeira auditoria)
)
PARAM['ME'] = PARAM['MG'] * PARAM['pct_ME']; PARAM['IT'] = PARAM['ME']; PARAM['IE'] = PARAM['IT'] * PARAM['pct_IE']

# Ilustracao 8.7 (Gramling) - fatores de confiabilidade, expansao e ajuste incremental (5 paginas do material da disciplina)
TAB87 = pd.DataFrame({
    'RD': [0.01, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.50],
    'FC': [4.61, 3.00, 2.31, 1.90, 1.61, 1.39, 1.21, 0.70],
    'FEE': [1.90, 1.60, 1.50, 1.40, 1.30, 1.25, 1.20, 1.00]})
INCR = {0.01: [1.03, .77, .64, .56, .50, .46, .43, .41, .38, .38], 0.05: [.75, .55, .46, .40, .36, .33, .30, .29, .27, .26],
        0.10: [.58, .44, .36, .31, .28, .26, .24, .22, .21, .20], 0.15: [.48, .34, .30, .25, .23, .21, .19, .18, .17, .17],
        0.20: [.39, .28, .24, .21, .18, .17, .16, .14, .14, .14], 0.25: [.31, .23, .18, .17, .15, .13, .13, .12, .11, .10],
        0.30: [.23, .18, .15, .13, .11, .11, .10, .09, .08, .08], 0.50: [0] * 10}

def ler_base():
    return pd.read_csv(f'{B}/BASE_IMOBILIZADO_RECV3_SINTETICA.csv.gz', sep=';', decimal=',', low_memory=False,
                       parse_dates=['data_aquisicao', 'data_inicio_depreciacao', 'data_baixa', 'data_cadastro_sistema', 'data_ultima_alteracao'],
                       dtype={'usuario_aprovacao': str, 'poco_id': str})

base = ler_base()
base['usuario_aprovacao'] = base.usuario_aprovacao.fillna('')
campos = pd.read_csv(f'{B}/PARAMETROS_CAMPOS_SIMULADOS.csv', sep=';', decimal=',', encoding='utf-8-sig')
print('linhas lidas:', len(base))

# =============================================================================
# (1) CONCILIACAO - base x NE 7.1 (F01 p.68-69) x CVM (F03, conta 1.02.03)
# =============================================================================
NE7_CTR = {'MAQUINAS_EQUIPAMENTOS': (716487, 103733), 'IMOBILIZADO_EM_ANDAMENTO': (68419, 0), 'DIREITO_PRODUCAO_OLEO_GAS': (2851955, 850980),
           'DESENVOLVIMENTO_CAMPOS': (4018461, 1749489), 'BLOCOS_EXPLORATORIOS': (9553, 0), 'ABANDONO_POCO_ARO': (73827, 36676),
           'ALMOXARIFADO_INVERSOES_FIXAS': (435242, 0), 'ADIANTAMENTOS': (15677, 0), 'OUTROS': (195997, 39032)}
NE7_CON = {'MAQUINAS_EQUIPAMENTOS': (718521, 103788, 38884), 'IMOBILIZADO_EM_ANDAMENTO': (72171, 0, 0), 'DIREITO_PRODUCAO_OLEO_GAS': (2931329, 922414, 132263),
           'DESENVOLVIMENTO_CAMPOS': (5186676, 2422704, 485386), 'BLOCOS_EXPLORATORIOS': (20157, 0, 0), 'ABANDONO_POCO_ARO': (79303, 41848, 1174),
           'ALMOXARIFADO_INVERSOES_FIXAS': (477651, 0, 0), 'ADIANTAMENTOS': (18419, 0, 0), 'OUTROS': (199325, 41957, 14161)}
g = base.groupby(['entidade', 'classe_ne7'])[['custo_historico', 'depreciacao_acumulada', 'valor_contabil_liquido', 'depreciacao_2025']].sum()
conc = []
for classe, (c, d, dep) in NE7_CON.items():
    pr = g.loc[('PETRORECONCAVO', classe)] if ('PETRORECONCAVO', classe) in g.index else pd.Series(0, index=g.columns)
    ti = g.loc[('SPE_TIETA', classe)] if ('SPE_TIETA', classe) in g.index else pd.Series(0, index=g.columns)
    conc.append(dict(classe_ne7=classe, n_registros=int((base.classe_ne7 == classe).sum()),
                     custo_base_ctr=pr.custo_historico, custo_ne7_ctr=NE7_CTR[classe][0] * 1000.0, custo_base_cons=pr.custo_historico + ti.custo_historico, custo_ne7_cons=c * 1000.0,
                     acum_base_ctr=pr.depreciacao_acumulada, acum_ne7_ctr=NE7_CTR[classe][1] * 1000.0, acum_base_cons=pr.depreciacao_acumulada + ti.depreciacao_acumulada, acum_ne7_cons=d * 1000.0,
                     dep2025_base_cons=pr.depreciacao_2025 + ti.depreciacao_2025, dep2025_ne7_cons=dep * 1000.0))
conc = pd.DataFrame(conc)
balancete = base.groupby(['conta_custo', 'descricao_conta'])[['custo_historico', 'depreciacao_acumulada', 'valor_contabil_liquido']].sum().reset_index()
balancete['n_registros'] = base.groupby('conta_custo').size().values

# =============================================================================
# (2) TESTES NA POPULACAO INTEIRA (100% dos registros - nao e amostragem, NBC TA 500 A52-A56 / 530 item 5)
# =============================================================================
exc = []
def registrar(teste, df, detalhe, valor=None):
    v = df.valor_contabil_liquido if valor is None else valor
    exc.append(pd.DataFrame({'teste': teste, 'id_registro': df.id_registro.values, 'classe_ne7': df.classe_ne7.values,
                             'valor_contabil_liquido': df.valor_contabil_liquido.values, 'valor_em_risco': np.asarray(v, float), 'detalhe': detalhe}))

# T01 duplicidade: mesma chave (documento, fornecedor, custo e data de aquisicao); excecao = registro cadastrado depois
chave = ['documento_origem', 'fornecedor_id', 'custo_historico', 'data_aquisicao']
d = base[base.duplicated(chave, keep=False)].sort_values(chave + ['data_cadastro_sistema', 'id_registro'])
d1 = d[d.duplicated(chave, keep='first')]
registrar('T01_DUPLICIDADE', d1, 'Mesmo documento/fornecedor/valor/data de outro registro (cadastro posterior)')
# T02 corte: data de aquisicao posterior a data-base
registrar('T02_CORTE', base[base.data_aquisicao > DATA_BASE], 'Data de aquisicao posterior a 31/12/2025')
# T03 valor liquido negativo
t3 = base[base.valor_contabil_liquido < 0]; registrar('T03_VCL_NEGATIVO', t3, 'Depreciacao acumulada maior que o custo', valor=t3.valor_contabil_liquido)
# T04 recalculo da depreciacao linear (taxa x tempo desde o inicio da depreciacao; tolerancia = maior entre 1% do custo e R$ 100)
L = base[base.metodo_depreciacao == 'LINEAR']
anos = ((DATA_BASE - L.data_inicio_depreciacao).dt.days / 365.25).clip(lower=0)
esperada = (L.custo_historico * np.minimum(1, anos * L.taxa_anual_linear)).round(2)
dif = L.depreciacao_acumulada - esperada
f4 = abs(dif) > np.maximum(0.01 * L.custo_historico, 100)
registrar('T04_RECALCULO_DEPRECIACAO', L[f4], 'Acumulada difere do recalculo (taxa x tempo)', valor=-dif[f4])
# T05 item baixado ainda com saldo
registrar('T05_BAIXADO_COM_SALDO', base[(base.status == 'BAIXADO') & (base.valor_contabil_liquido != 0)], 'Status BAIXADO com saldo remanescente')
# T06 almoxarifado sem movimentacao ha mais de 5 anos
registrar('T06_ALMOX_SEM_MOVIMENTO_5A', base[(base.classe_ne7 == 'ALMOXARIFADO_INVERSOES_FIXAS') & (base.data_ultima_alteracao < pd.Timestamp('2021-01-01'))],
          'Sem movimentacao desde antes de 01/01/2021 (avaliar obsolescencia / valor recuperavel)')
# T07 segregacao de funcoes no cadastro (maker-checker)
registrar('T07_MAKER_CHECKER', base[(base.usuario_aprovacao == '') | (base.usuario_aprovacao == base.usuario_cadastro)],
          'Cadastro sem aprovacao ou aprovado pelo proprio cadastrador', valor=0)
# T08 campo inexistente no cadastro de concessoes (existencia / classificacao)
validos = set(campos.campo) | {'UPGN GUAMARE'} | set(base.loc[base.classe_ne7 == 'BLOCOS_EXPLORATORIOS', 'campo'])
registrar('T08_CAMPO_NAO_CADASTRADO', base[~base.campo.isin(validos)], 'Campo/local nao consta das 58 concessoes, blocos ou UPGN')
# T09 imobilizado em andamento parado ha mais de 2 anos
registrar('T09_ANDAMENTO_PARADO_2A', base[(base.classe_ne7 == 'IMOBILIZADO_EM_ANDAMENTO') & (base.data_ultima_alteracao < pd.Timestamp('2024-01-01'))],
          'Ordem em andamento sem movimentacao desde antes de 01/01/2024 (transferencia tardia / impairment)')
# T10 data de aquisicao anterior a constituicao da companhia (21/07/1999) em capex proprio
registrar('T10_DATA_ANTERIOR_CONSTITUICAO', base[(base.origem == 'CAPEX_PROPRIO') & (base.data_aquisicao < pd.Timestamp('1999-07-21'))],
          'Data de aquisicao anterior a 21/07/1999', valor=0)
excecoes = pd.concat(exc, ignore_index=True)
excecoes.to_csv(f'{RES}/EXCECOES_TESTES_POPULACAO.csv', index=False, sep=';', decimal=',', encoding='utf-8-sig')
resumo_testes = excecoes.groupby('teste').agg(qtd=('id_registro', 'count'), valor_em_risco=('valor_em_risco', 'sum')).reset_index()
print(resumo_testes)

# =============================================================================
# (3) MUS - Amostragem de unidades monetarias (Gramling, cap. 8; NBC TA 530)
# Afirmacoes: EXISTENCIA e AVALIACAO (teste de superavaliacao). Unidade amostral = R$ 1 do valor contabil liquido.
# Populacao: itens com VCL > 0. Itens com VCL <= 0 sao tratados a parte (T03), como recomenda o material.
# =============================================================================
RD = PARAM['RA'] / (PARAM['RI'] * PARAM['RC'])
col = TAB87[TAB87.RD <= RD + 1e-12].RD.max()          # se o RD calculado nao estiver na tabela, usa a coluna imediatamente inferior (conservador)
FC = float(TAB87.loc[TAB87.RD == col, 'FC'].iloc[0]); FEE = float(TAB87.loc[TAB87.RD == col, 'FEE'].iloc[0]); FATORES = INCR[col]
I = (PARAM['IT'] - PARAM['IE'] * FEE) / FC                                # intervalo de amostragem
pop = base[base.valor_contabil_liquido > 0].copy().sort_values('id_registro').reset_index(drop=True)   # ordem da populacao = id do registro
VC_POP = pop.valor_contabil_liquido.sum()
n_max = VC_POP / I
rng = np.random.default_rng(SEED_MUS); inicio = float(rng.uniform(1, I))  # ponto de partida aleatorio entre 1 e I
pop['acumulado'] = pop.valor_contabil_liquido.cumsum()
pontos = inicio + I * np.arange(int(np.floor((VC_POP - inicio) / I)) + 1)
pos = np.searchsorted(pop.acumulado.values, pontos)                        # item cujo acumulado alcanca o ponto
sel = pd.DataFrame({'ponto_selecao': pontos, 'pos': pos}).groupby('pos').agg(pontos_no_item=('ponto_selecao', 'count'), primeiro_ponto=('ponto_selecao', 'min')).reset_index()
amostra = pop.iloc[sel.pos].reset_index(drop=True).join(sel[['pontos_no_item', 'primeiro_ponto']])
amostra['estrato'] = np.where(amostra.valor_contabil_liquido >= I, 'SUPERIOR (VC >= I)', 'INFERIOR')
# Inspecao documental SIMULADA somente dos itens selecionados
insp = pd.read_csv(f'{B}/INSPECAO_DOCUMENTAL_SIMULADA.csv.gz', sep=';', decimal=',')
amostra = amostra.merge(insp[['id_registro', 'documento_localizado', 'data_documento', 'valor_liquido_suportado', 'observacao_inspecao']], on='id_registro', how='left')
amostra['valor_auditado'] = amostra.valor_liquido_suportado
amostra['incorrecao'] = (amostra.valor_contabil_liquido - amostra.valor_auditado).round(2)
amostra['tainting'] = np.where(amostra.estrato == 'INFERIOR', amostra.incorrecao / amostra.valor_contabil_liquido, np.nan)
amostra['incorrecao_projetada'] = np.where(amostra.estrato == 'INFERIOR', amostra.tainting * I, amostra.incorrecao)
erros = amostra[amostra.incorrecao > 0.005]
inf = erros[erros.estrato == 'INFERIOR'].sort_values('tainting', ascending=False).reset_index(drop=True)
fat = [FATORES[i] if i < len(FATORES) else FATORES[-1] for i in range(len(inf))]   # >10 erros: usa o ultimo fator da tabela (premissa documentada)
inf['fator_incremental'] = fat; inf['ajuste_incremental'] = inf.tainting * inf.fator_incremental * I
PB = FC * I
IMP = erros.incorrecao_projetada.sum()
IA = inf.ajuste_incremental.sum()
LSI = PB + IMP + IA
sub = amostra[amostra.incorrecao < -0.005]
MUS = dict(RD=RD, coluna_tabela=col, FC=FC, FEE=FEE, IT=PARAM['IT'], IE=PARAM['IE'], I=I, VC_populacao=VC_POP, n_populacao=len(pop), n_max=n_max,
           ponto_partida=inicio, n_pontos=len(pontos), n_itens_distintos=len(amostra), n_estrato_superior=int((amostra.estrato != 'INFERIOR').sum()),
           valor_amostra=amostra.valor_contabil_liquido.sum(), n_erros=len(erros), n_erros_inferior=len(inf), PB=PB, IMP=IMP, IA=IA, LSI=LSI,
           conclusao='ACEITAR (LSI <= IT)' if LSI <= PARAM['IT'] else 'NAO ACEITAR SEM TRABALHO ADICIONAL (LSI > IT)', n_subavaliacoes=len(sub))
print(json.dumps({k: (round(v, 2) if isinstance(v, float) else v) for k, v in MUS.items()}, ensure_ascii=False))

# =============================================================================
# (4) COMPARACAO - erros inseridos x detectados (o LOG so e lido AGORA)
# =============================================================================
lg = pd.read_csv(f'{B}/controle_erros/LOG_ERROS_INSERIDOS.csv', sep=';', decimal=',', encoding='utf-8-sig')
MAPA = {'E01': 'T01_DUPLICIDADE', 'E02': 'T02_CORTE', 'E03': 'T03_VCL_NEGATIVO', 'E04': 'T04_RECALCULO_DEPRECIACAO', 'E05': 'T05_BAIXADO_COM_SALDO',
        'E06': 'T06_ALMOX_SEM_MOVIMENTO_5A', 'E07': 'T07_MAKER_CHECKER', 'E08': 'T08_CAMPO_NAO_CADASTRADO', 'E09': 'MUS (inspecao documental)',
        'E10': 'T09_ANDAMENTO_PARADO_2A', 'E11': 'T10_DATA_ANTERIOR_CONSTITUICAO'}
comp = []
for tipo, teste in MAPA.items():
    ins = set(lg[lg.tipo_erro == tipo].id_registro)
    det = set(amostra[amostra.incorrecao > 0.005].id_registro) if tipo == 'E09' else set(excecoes[excecoes.teste == teste].id_registro)
    vp = ins & det
    outros_tipos = det - ins
    comp.append(dict(tipo_erro=tipo, descricao=lg[lg.tipo_erro == tipo].descricao.iloc[0], teste=teste, inseridos=len(ins), detectados_pelo_teste=len(det),
                     verdadeiros_positivos=len(vp), nao_detectados=len(ins - det),
                     detectados_de_outro_tipo_inserido=len(outros_tipos & set(lg.id_registro)), falsos_positivos=len(outros_tipos - set(lg.id_registro)),
                     taxa_deteccao=len(vp) / len(ins) if ins else np.nan, efeito_inserido=lg[lg.tipo_erro == tipo]['efeito_monetario_R$'].sum()))
comp = pd.DataFrame(comp)
amostra['erro_inserido'] = amostra.id_registro.map(lg.groupby('id_registro').tipo_erro.apply(lambda s: ','.join(sorted(set(s)))))
total_inserido_superaval = lg.loc[lg['efeito_monetario_R$'] > 0, 'efeito_monetario_R$'].sum()
VAL = dict(incorrecao_projetada_MUS=IMP, LSI=LSI, superavaliacao_total_inserida=total_inserido_superaval,
           LSI_cobre_erro_real='SIM' if LSI >= total_inserido_superaval else 'NAO')
print(comp[['tipo_erro', 'inseridos', 'verdadeiros_positivos', 'falsos_positivos', 'taxa_deteccao']]); print(VAL)

# =============================================================================
# Saidas: planilhas editaveis (com formulas) e CSVs
# =============================================================================
HEAD = PatternFill('solid', fgColor='1F4E78'); thin = Side(style='thin', color='BFBFBF'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
SINT = PatternFill('solid', fgColor='FFF2CC')
RODAPE = 'DADOS SINTETICOS - procedimento SIMULADO PARA FINS DIDATICOS | Preparador: matricula 211059370 | Data: 23/10/2026 | Revisao: ______ (nao preenchido)'
def cab(ws, r, cols, w=None):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = Font(bold=True, color='FFFFFF'); x.fill = HEAD; x.alignment = Alignment(wrap_text=True); x.border = BD
    if w:
        for j, k in enumerate(w, 1): ws.column_dimensions[get_column_letter(j)].width = k
def tit(ws, t):
    ws['A1'] = t; ws['A1'].font = Font(bold=True, size=13); ws['A2'] = RODAPE; ws['A2'].font = Font(italic=True, size=9, color='9C5700'); ws['A2'].fill = SINT
def df_to(ws, df, r0, nf=None):
    for i, row in enumerate(df.itertuples(index=False), r0):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)) and np.isnan(v): v = None
            elif isinstance(v, np.generic): v = v.item()
            elif isinstance(v, pd.Timestamp): v = v.date()
            c = ws.cell(row=i, column=j, value=v); c.border = BD
            if nf and j in nf: c.number_format = nf[j]
    return r0 + len(df) - 1
NUM = '#,##0.00;(#,##0.00);-'; PCT = '0.00%'

# ---- CONC_BASE_RECV3.xlsx
wb = Workbook(); ws = wb.active; ws.title = 'Conc_classe'
tit(ws, 'Conciliacao: base simulada x NE 7.1 (F01 p.68-69) - R$ (centavos)')
cols = ['Classe NE 7', 'Registros', 'Custo base (ctr)', 'Custo NE 7 (ctr)', 'Dif.', 'Custo base (cons)', 'Custo NE 7 (cons)', 'Dif.',
        'Acum. base (ctr)', 'Acum. NE 7 (ctr)', 'Dif.', 'Acum. base (cons)', 'Acum. NE 7 (cons)', 'Dif.', 'Deprec. 2025 base (cons)', 'Deprec. 2025 NE 7 (cons)', 'Dif.']
cab(ws, 4, cols, [32, 10] + [17] * 15)
for i, r in enumerate(conc.itertuples(index=False), 5):
    vals = [r.classe_ne7, r.n_registros, r.custo_base_ctr, r.custo_ne7_ctr, f'=C{i}-D{i}', r.custo_base_cons, r.custo_ne7_cons, f'=F{i}-G{i}',
            r.acum_base_ctr, r.acum_ne7_ctr, f'=I{i}-J{i}', r.acum_base_cons, r.acum_ne7_cons, f'=L{i}-M{i}', r.dep2025_base_cons, r.dep2025_ne7_cons, f'=O{i}-P{i}']
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BD; c.number_format = NUM if j > 2 else 'General'
t = 5 + len(conc)
ws.cell(row=t, column=1, value='TOTAL IMOBILIZADO').font = Font(bold=True)
for j in range(2, 18):
    L_ = get_column_letter(j); ws.cell(row=t, column=j, value=f'=SUM({L_}5:{L_}{t-1})').number_format = NUM
ws.cell(row=t + 1, column=1, value='Valor liquido consolidado - base'); ws.cell(row=t + 1, column=6, value=f'=F{t}-L{t}').number_format = NUM
ws.cell(row=t + 2, column=1, value='Valor liquido consolidado - CVM 1.02.03 (F03) / BP (F01 p.51, R$ mil x 1.000)'); ws.cell(row=t + 2, column=6, value=6170841000).number_format = NUM
ws.cell(row=t + 3, column=1, value='Diferenca (arredondamento da DF em R$ mil)'); ws.cell(row=t + 3, column=6, value=f'=F{t+1}-F{t+2}').number_format = NUM
notas = ['Explicacao das diferencas: (1) A DF e publicada em R$ mil (arredondada); a base, em centavos. A geracao introduz, por classe/entidade, diferenca simulada de',
         'arredondamento de no maximo R$ 240, de modo que cada subtotal arredondado para R$ mil coincide com a NE 7.1. (2) A depreciacao do exercicio dos itens LINEARES',
         'resulta das datas simuladas e excede a NE 7.1 (a base foi calibrada para os SALDOS e para a deplecao UoP); a diferenca nao afeta a conciliacao do saldo e e uma',
         'limitacao do dado sintetico. (3) Intangivel (software, R$ 83.384 mil) fora da base: a conta escolhida e o imobilizado (98,7% da rubrica Imobilizado e intangivel).',
         '(4) Os erros inseridos estao DENTRO do saldo contabil (a base reproduz um razao que contem distorcoes simuladas); por isso o total da base continua igual a DF.']
for k, n in enumerate(notas): ws.cell(row=t + 5 + k, column=1, value=n)
ws2 = wb.create_sheet('Balancete_plano_simulado'); tit(ws2, 'Balancete do imobilizado por conta do PLANO DE CONTAS SIMULADO')
cab(ws2, 4, ['Conta', 'Descricao (simulada)', 'Custo', 'Deprec. acumulada', 'Valor liquido', 'Registros'], [16, 70, 18, 18, 18, 10])
e = df_to(ws2, balancete, 5, {3: NUM, 4: NUM, 5: NUM})
ws2.cell(row=e + 1, column=2, value='TOTAL').font = Font(bold=True)
for j in (3, 4, 5, 6): ws2.cell(row=e + 1, column=j, value=f'=SUM({get_column_letter(j)}5:{get_column_letter(j)}{e})').number_format = NUM
wb.save('planilhas/CONC_BASE_RECV3.xlsx')

# ---- TESTES_BASE_RECV3.xlsx (resumo dos testes de populacao, MUS e comparacao)
wb = Workbook(); ws = wb.active; ws.title = 'Resumo_testes_populacao'
tit(ws, 'Testes em 100% da populacao (T01-T10) - resumo')
cab(ws, 4, ['Teste', 'Qtd. excecoes', 'Valor em risco (R$)', 'Afirmacao', 'Norma'], [34, 14, 20, 30, 26])
AFIRM = {'T01_DUPLICIDADE': ('Existencia', 'NBC TA 500 / 240'), 'T02_CORTE': ('Existencia / corte', 'NBC TA 500 / 330'), 'T03_VCL_NEGATIVO': ('Avaliacao', 'NBC TA 500 / 520'),
         'T04_RECALCULO_DEPRECIACAO': ('Avaliacao (exatidao)', 'NBC TA 500 A21 (recalculo)'), 'T05_BAIXADO_COM_SALDO': ('Existencia', 'NBC TA 500'),
         'T06_ALMOX_SEM_MOVIMENTO_5A': ('Avaliacao (recuperabilidade)', 'NBC TA 540 / 501'), 'T07_MAKER_CHECKER': ('Controle - segregacao', 'NBC TA 315 / 265'),
         'T08_CAMPO_NAO_CADASTRADO': ('Existencia / classificacao', 'NBC TA 500 / 240'), 'T09_ANDAMENTO_PARADO_2A': ('Avaliacao / classificacao', 'NBC TA 540'),
         'T10_DATA_ANTERIOR_CONSTITUICAO': ('Integridade cadastral', 'NBC TA 500 (confiabilidade dos dados)')}
for i, r in enumerate(resumo_testes.itertuples(index=False), 5):
    for j, v in enumerate([r.teste, int(r.qtd), float(r.valor_em_risco), *AFIRM[r.teste]], 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BD; c.number_format = NUM if j == 3 else 'General'
ws3 = wb.create_sheet('Excecoes_detalhe'); tit(ws3, 'Detalhe das excecoes (todas)')
cab(ws3, 4, list(excecoes.columns), [30, 16, 30, 18, 18, 70]); df_to(ws3, excecoes, 5, {4: NUM, 5: NUM})

# MUS - parametros com formulas
wm = wb.create_sheet('MUS_parametros'); tit(wm, 'MUS - parametros e intervalo (Gramling cap. 8; NBC TA 530) - formulas editaveis')
cab(wm, 4, ['Parametro', 'Valor', 'Fonte / formula'], [52, 22, 80])
linhas = [('Materialidade global provisoria - MG (R$)', PARAM['MG'], 'PT 320: 1% da receita liquida consolidada 2025 (R$ 3.157.609 mil) = 31.576 mil, arredondado p/ baixo (NBC TA 320 itens 10, A5, A8)'),
          ('% da MG para execucao', PARAM['pct_ME'], 'Julgamento (NBC TA 320 item 11 e A13): primeira auditoria + inconsistencias de divulgacao + risco de impairment'),
          ('Materialidade para execucao - ME (R$)', '=B5*B6', 'MG x %'),
          ('Incorrecao toleravel - IT (R$)', '=B7', 'IT = ME (sem materialidade especifica para o imobilizado)'),
          ('% IT para incorrecao esperada', PARAM['pct_IE'], 'Premissa do auditor (sem historico: primeira auditoria)'),
          ('Incorrecao esperada - IE (R$)', '=B8*B9', 'IT x %'),
          ('Risco de auditoria - RA', PARAM['RA'], 'Gramling: RA = RI x RC x RD'), ('Risco inerente - RI', PARAM['RI'], 'Maximo: estimativas e julgamentos (PT 315)'),
          ('Risco de controle - RC', PARAM['RC'], 'Maximo: controles nao testados (sem acesso)'), ('Risco de deteccao - RD', '=B11/(B12*B13)', 'RD = RA / (RI x RC)'),
          ('Coluna da Ilustracao 8.7 utilizada', col, 'RD da tabela imediatamente inferior ou igual ao calculado'),
          ('Fator de confiabilidade - FC', FC, 'Ilustracao 8.7'), ('Fator de expansao de erro - FEE', FEE, 'Ilustracao 8.7'),
          ('Intervalo de amostragem - I (R$)', '=(B8-B10*B17)/B16', 'I = (IT - IE x FEE) / FC'),
          ('Valor contabil da populacao (VCL > 0) (R$)', VC_POP, f'{len(pop):,} itens; itens com VCL <= 0 testados a parte (T03)'.replace(',', '.')),
          ('Tamanho maximo da amostra - n', '=B19/B18', 'n = valor contabil / I'),
          ('Ponto de partida aleatorio (R$)', inicio, f'Gerador numpy, semente {SEED_MUS} (entre 1 e I)'),
          ('Pontos de selecao', len(pontos), 'inicio + k x I'), ('Itens distintos selecionados', len(amostra), 'itens com mais de um ponto contam uma vez'),
          ('Itens do estrato superior (VC >= I)', MUS['n_estrato_superior'], 'selecao certa (100%)')]
for i, (a, v, f) in enumerate(linhas, 5):
    wm.cell(row=i, column=1, value=a); c = wm.cell(row=i, column=2, value=v); wm.cell(row=i, column=3, value=f)
    c.number_format = '0.00%' if i in (6, 9, 11, 12, 13, 14, 15) else ('0.00' if i in (16, 17, 20) else NUM)
    for j in (1, 2, 3): wm.cell(row=i, column=j).border = BD
wm.cell(row=27, column=1, value='Ilustracao 8.7 (Gramling) - fatores').font = Font(bold=True)
cab(wm, 28, ['Risco de deteccao', 'FC', 'FEE'] + [f'Incr. {k}' for k in range(1, 11)])
for i, rd in enumerate(TAB87.RD, 29):
    for j, v in enumerate([rd, float(TAB87.FC[i - 29]), float(TAB87.FEE[i - 29])] + INCR[rd], 1): wm.cell(row=i, column=j, value=v).border = BD

# MUS - selecao e resultado da inspecao
wa = wb.create_sheet('MUS_amostra'); tit(wa, 'MUS - itens selecionados e resultado da inspecao documental SIMULADA')
cols_a = ['id_registro', 'classe_ne7', 'conta_custo', 'campo', 'tipo_item', 'estrato', 'pontos_no_item', 'primeiro_ponto', 'acumulado', 'valor_contabil_liquido',
          'documento_localizado', 'data_documento', 'valor_auditado', 'incorrecao', 'tainting', 'incorrecao_projetada', 'observacao_inspecao']
cab(wa, 4, cols_a, [16, 28, 14, 22, 30, 18, 8, 16, 18, 16, 8, 12, 16, 14, 10, 16, 50])
I_REF = "MUS_parametros!$B$18"
for i, r in enumerate(amostra[cols_a].itertuples(index=False), 5):
    for j, v in enumerate(r, 1):
        if isinstance(v, (np.floating, float)) and np.isnan(v): v = None
        elif isinstance(v, np.generic): v = v.item()
        c = wa.cell(row=i, column=j, value=v); c.border = BD
        if j in (8, 9, 10, 13, 14, 16): c.number_format = NUM
    wa.cell(row=i, column=14, value=f'=J{i}-M{i}').number_format = NUM
    wa.cell(row=i, column=15, value=f'=IF(F{i}="INFERIOR",N{i}/J{i},"")').number_format = PCT
    wa.cell(row=i, column=16, value=f'=IF(F{i}="INFERIOR",O{i}*{I_REF},N{i})').number_format = NUM
last_a = 4 + len(amostra)

# MUS - avaliacao com formulas (LSI)
wv = wb.create_sheet('MUS_avaliacao'); tit(wv, 'MUS - avaliacao: LSI = precisao basica + incorrecao mais provavel + ajuste incremental')
cab(wv, 4, ['Componente', 'Valor (R$)', 'Formula / explicacao'], [52, 22, 90])
aval = [('Precisao basica (PB)', f'=MUS_parametros!B16*{I_REF}', 'FC x I (LSI se nenhum erro fosse encontrado)'),
        ('Incorrecao mais provavel - estrato inferior (projetada)', f'=SUMIFS(MUS_amostra!P5:P{last_a},MUS_amostra!F5:F{last_a},"INFERIOR",MUS_amostra!N5:N{last_a},">0")', 'soma de tainting x I'),
        ('Incorrecao - estrato superior (real, sem projecao)', f'=SUMIFS(MUS_amostra!N5:N{last_a},MUS_amostra!F5:F{last_a},"<>INFERIOR",MUS_amostra!N5:N{last_a},">0")', 'diferenca efetiva'),
        ('Incorrecao mais provavel (IMP)', '=B6+B7', 'projetada + real'),
        ('Ajuste incremental (IA)', '=SUM(E13:E200)', 'taintings em ordem decrescente x fator incremental x I (tabela abaixo)'),
        ('LIMITE SUPERIOR DE INCORRECAO (LSI)', '=B5+B8+B9', 'PB + IMP + IA'),
        ('Incorrecao toleravel (IT)', '=MUS_parametros!B8', ''), ('Conclusao', '=IF(B10<=B11,"ACEITAR: LSI <= IT","NAO ACEITAR SEM TRABALHO ADICIONAL: LSI > IT")', 'regra de decisao do material (pag. 302)')]
for i, (a, f, x) in enumerate(aval, 5):
    wv.cell(row=i, column=1, value=a).border = BD; c = wv.cell(row=i, column=2, value=f); c.border = BD; c.number_format = NUM; wv.cell(row=i, column=3, value=x).border = BD
wv['B12'].number_format = 'General'; wv['A10'].font = Font(bold=True); wv['B10'].font = Font(bold=True)
cab(wv, 12, ['Ordem', 'id_registro', 'Tainting', 'Fator incremental (col. RD)', 'Ajuste incremental (R$)'])
for k, r in enumerate(inf.itertuples(index=False), 13):
    wv.cell(row=k, column=1, value=k - 12); wv.cell(row=k, column=2, value=r.id_registro)
    wv.cell(row=k, column=3, value=float(r.tainting)).number_format = PCT; wv.cell(row=k, column=4, value=float(r.fator_incremental))
    wv.cell(row=k, column=5, value=f'=C{k}*D{k}*{I_REF}').number_format = NUM
nota_row = 14 + len(inf)
for k, n in enumerate(['Premissa: a Ilustracao 8.7 traz fatores ate o 10o erro; para erros adicionais foi usado o ultimo fator da coluna (premissa conservadora documentada).',
                       'Subavaliacoes (valor auditado > contabil) nao entram no LSI de superavaliacao; sao avaliadas a parte no PT 450.',
                       'Resultado calculado em Python (conferencia): ' + f'PB={PB:,.2f}; IMP={IMP:,.2f}; IA={IA:,.2f}; LSI={LSI:,.2f}; IT={PARAM["IT"]:,.2f}.']):
    wv.cell(row=nota_row + k, column=1, value=n)

wc = wb.create_sheet('Inseridos_x_detectados'); tit(wc, 'Comparacao: erros INSERIDOS (log de controle) x DETECTADOS pelos testes')
cab(wc, 4, list(comp.columns), [8, 60, 30, 10, 12, 12, 12, 14, 12, 12, 18])
df_to(wc, comp, 5, {10: '0.0%', 11: NUM})
r_ = 7 + len(comp)
for k, n in enumerate([f'MUS: incorrecao mais provavel projetada = R$ {IMP:,.2f}; LSI = R$ {LSI:,.2f}.',
                       f'Superavaliacao total efetivamente inserida na base (todos os tipos monetarios) = R$ {total_inserido_superaval:,.2f}.',
                       f'O LSI {"cobre" if LSI >= total_inserido_superaval else "NAO cobre"} a superavaliacao real inserida - validacao didatica do metodo.',
                       'Os testes T01-T10 foram executados antes da leitura do log (as cegas); o log foi usado apenas nesta aba.']):
    wc.cell(row=r_ + k, column=1, value=n)
wb.save('planilhas/TESTES_BASE_RECV3.xlsx')

amostra.to_csv(f'{RES}/MUS_AMOSTRA_SELECIONADA.csv', index=False, sep=';', decimal=',', encoding='utf-8-sig')
comp.to_csv(f'{RES}/COMPARATIVO_INSERIDOS_DETECTADOS.csv', index=False, sep=';', decimal=',', encoding='utf-8-sig')
json.dump({'parametros': PARAM, 'mus': MUS, 'validacao': VAL}, open(f'{RES}/RESULTADOS_MUS.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=float)
print('ok')
