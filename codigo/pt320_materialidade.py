# -*- coding: utf-8 -*-
# PT 320 - Materialidade (NBC TA 320 (R1); NBC TA 450 (R1) item 5 e A2) - Preparador: matricula 211059370
# Gera planilhas/PT320_MATERIALIDADE_RECV3.xlsx com FORMULAS (valores de entrada = DF consolidada 2025/2024, F01 p.51-56).
# Uso: python codigo/pt320_materialidade.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook(); ws = wb.active; ws.title = 'PT320_calculo'
H = PatternFill('solid', fgColor='1F4E78'); thin = Side(style='thin', color='BFBFBF'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);-'; PCT = '0.0%'
ws['A1'] = 'PT 320 - Materialidade no planejamento e na execucao (NBC TA 320 (R1)) - R$ mil'; ws['A1'].font = Font(bold=True, size=13)
ws['A2'] = 'PetroReconcavo S.A. - DF consolidadas 31/12/2025 | Preparador: matricula 211059370 | 24/10/2026 | Revisao: ______ (nao preenchido)'
for j, w in enumerate([58, 16, 16, 12, 62], 1): ws.column_dimensions[get_column_letter(j)].width = w
def cab(r, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=j, value=c); x.font = Font(bold=True, color='FFFFFF'); x.fill = H; x.border = BD; x.alignment = Alignment(wrap_text=True)
def lin(r, vals, nfs=None):
    for j, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=j, value=v); x.border = BD
        if nfs and j in nfs: x.number_format = nfs[j]

cab(4, ['1. Dados dos referenciais (F01; F03)', '2025', '2024', 'Var. %', 'Fonte'])
dados = [('Receita liquida', 3157609, 3264554, 'DRE consolidada, F01 p.52'), ('Lucro antes dos impostos (LAIR)', 776194, 363405, 'DRE consolidada, F01 p.52'),
         ('  dos quais: resultado de instrumentos financeiros (swaps e collar)', 366796, -368547, 'NE 20, F01 p.94'),
         ('Lucro bruto', 1021436, 1191749, 'DRE consolidada, F01 p.52'), ('EBITDA (lucro operacional + D&A)', 1442656, 1643036, 'DIAG_RECV3.xlsx, aba Indicadores'),
         ('Ativo total', 8668015, 7437571, 'BP consolidado, F01 p.51'), ('Patrimonio liquido', 4330716, 4235277, 'BP consolidado, F01 p.51')]
for i, (n, a, b, f) in enumerate(dados, 5):
    lin(i, [n, a, b, f'=IFERROR(B{i}/C{i}-1,"")', f], {2: NUM, 3: NUM, 4: PCT})
# r5 receita, r6 LAIR, r7 swaps, r8 LB, r9 EBITDA, r10 AT, r11 PL
cab(13, ['2. Volatilidade do LAIR (item A5)', 'Valor', '', '', 'Leitura'])
lin(14, ['Variacao do LAIR 2025 x 2024', '=B6/C6-1', '', '', 'LAIR mais que dobrou por efeito de swaps -> referencial VOLATIL'], {2: PCT})
lin(15, ['Variacao da receita liquida 2025 x 2024', '=B5/C5-1', '', '', 'Receita estavel -> referencial adequado (A5: "receita total")'], {2: PCT})
cab(17, ['3. Referenciais testados (A5, A6, A8)', 'Base', '% (A8 / julgamento)', 'Resultado', 'Uso'])
ref = [('Receita liquida 2025 (ESCOLHIDO - A5)', '=B5', 0.01, 'Referencial escolhido: estavel; % do exemplo do A8 para receita total'),
       ('LAIR 2025 como divulgado', '=B6', 0.05, 'Nao usado: volatil (A5); serve so de comparacao'),
       ('LAIR normalizado - media 2024-2025 (A6)', '=(B6+C6)/2', 0.05, 'Teste de razoabilidade (A6: valor normalizado com base em resultados anteriores)'),
       ('LAIR 2025 sem o resultado de swaps (A6)', '=B6-B7', 0.05, 'Teste de razoabilidade (exclui evento excepcional)'),
       ('Ativo total 2025 (A4 d - entidade com divida relevante)', '=B10', 0.005, 'Teste de razoabilidade (% baixo; julgamento)')]
for i, (n, b, p, u) in enumerate(ref, 18):
    lin(i, [n, b, p, f'=B{i}*C{i}', u], {2: NUM, 3: '0.0%', 4: NUM})
cab(24, ['4. Determinacao (itens 10, 11; A13; NBC TA 450 item 5/A2)', 'Valor', '% / regra', '', 'Justificativa'])
lin(25, ['Materialidade calculada (1% da receita liquida)', '=D18', '', '', 'Item 10 e A5/A8'], {2: NUM})
lin(26, ['MATERIALIDADE GLOBAL - MG (arredondada para baixo)', '=FLOOR(B25,500)', '', '', 'Arredondamento conservador para multiplo de R$ 500 mil'], {2: NUM})
lin(27, ['% da MG para execucao (julgamento - A13)', 0.65, '', '', 'Primeira auditoria; inconsistencias de divulgacao ja identificadas; risco significativo de impairment'], {2: PCT})
lin(28, ['MATERIALIDADE PARA EXECUCAO - ME', '=B26*B27', '', '', 'Item 11 / item 9 (definicao)'], {2: NUM})
lin(29, ['% da MG para distorcoes claramente triviais (julgamento - NBC TA 450 A2)', 0.05, '', '', 'Magnitude "totalmente diferente (menor)" que a MG'], {2: PCT})
lin(30, ['LIMITE DE DISTORCAO CLARAMENTE TRIVIAL', '=B26*B29', '', '', 'NBC TA 450 item 5: acumular distorcoes acima deste valor'], {2: NUM})
cab(32, ['5. Materialidade especifica (item 10; A11-A12)', 'Valor', 'ME especifica', '', 'Fundamento'])
lin(33, ['Remuneracao da administracao e transacoes com partes relacionadas', 1500, '=B33*$B$27', '', 'A11 (1o topico): leis/regulamentos afetam expectativas dos usuarios; total da remuneracao R$ 30.694 mil (NE 16.2)'], {2: NUM, 3: NUM})
lin(34, ['Indices financeiros dos covenants (NE 9.3)', 'qualitativa', '-', '', 'Qualquer distorcao que altere o cumprimento de covenant e relevante independentemente do valor (A11; NBC TA 450 A21)'])
lin(35, ['Divulgacao da analise de sensibilidade da provisao para abandono (estimativa de alta incerteza)', 'qualitativa', '-', '', 'A11 (2o topico): analise de sensibilidade de estimativas com alta incerteza'])
cab(37, ['6. Reflexos nos outros PTs', 'Valor', '', '', 'Onde'])
lin(38, ['Incorrecao toleravel do MUS do imobilizado (IT = ME)', '=B28', '', '', 'PT 530 e codigo/testes_auditoria_base.py (PARAM)'], {2: NUM})
lin(39, ['Acumulacao de distorcoes (limite trivial)', '=B30', '', '', 'PT 450'], {2: NUM})
lin(40, ['Limite de investigacao dos procedimentos analiticos (<= ME)', '=B28', '', '', 'PT 520'], {2: NUM})
ws['A42'] = 'Item 12/13/A14: a MG foi calculada sobre valores REAIS do exercicio encerrado (DF publicada), nao sobre projecoes; revisar se surgirem informacoes que levariam a valor diferente.'
ws['A43'] = 'Classificacao: REALIZADO COM EVIDENCIA PUBLICA (calculo sobre a DF consolidada publicada).'
wb.save('planilhas/PT320_MATERIALIDADE_RECV3.xlsx'); print('ok PT320 xlsx')
