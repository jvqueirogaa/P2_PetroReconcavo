# -*- coding: utf-8 -*-
# P2 - Auditoria 1 - RECV3 - Preparador 211059370
# Gera o DICIONARIO DE VARIAVEIS (xlsx + md) e CONFERE se cada variavel documentada existe nos arquivos gerados.
# Uso: python codigo/gerar_dicionario.py
import os, gzip
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARQ = {
 'BASE_IMOBILIZADO_RECV3_SINTETICA.csv.gz': ('base_simulada/BASE_IMOBILIZADO_RECV3_SINTETICA.csv.gz', 'Base obrigatoria: subrazao analitico do imobilizado em 31/12/2025 - 1 linha = 1 item patrimonial/componente capitalizado'),
 'INSPECAO_DOCUMENTAL_SIMULADA.csv.gz': ('base_simulada/INSPECAO_DOCUMENTAL_SIMULADA.csv.gz', 'Resultado SIMULADO da inspecao de documentos (usado so para itens selecionados na amostra)'),
 'PARAMETROS_CAMPOS_SIMULADOS.csv': ('base_simulada/PARAMETROS_CAMPOS_SIMULADOS.csv', 'Cadastro das 58 concessoes com producao e reservas (rateio sintetico de totais reais)'),
 'PLANO_CONTAS_IMOBILIZADO_SIMULADO.csv': ('base_simulada/PLANO_CONTAS_IMOBILIZADO_SIMULADO.csv', 'Plano de contas analitico SIMULADO do imobilizado'),
 'CADASTRO_POCOS_SIMULADO.csv': ('base_simulada/CADASTRO_POCOS_SIMULADO.csv', 'Cadastro sintetico de pocos'),
 'LOG_ERROS_INSERIDOS.csv': ('base_simulada/controle_erros/LOG_ERROS_INSERIDOS.csv', 'Log de controle dos erros inseridos na base do imobilizado (nao usado pelos testes)'),
 'CADASTRO_EMPREGADOS_SIMULADO.csv': ('base_simulada/folha/CADASTRO_EMPREGADOS_SIMULADO.csv', 'Base complementar: cadastro de empregados'),
 'FOLHA_PAGAMENTOS_2025_SIMULADA.csv': ('base_simulada/folha/FOLHA_PAGAMENTOS_2025_SIMULADA.csv', 'Base complementar: lancamentos da folha de 2025'),
 'OBRIGACOES_TRABALHISTAS_31122025_SIMULADAS.csv': ('base_simulada/folha/OBRIGACOES_TRABALHISTAS_31122025_SIMULADAS.csv', 'Base complementar: obrigacoes trabalhistas em 31/12/2025'),
 'LOG_ERROS_INSERIDOS_FOLHA.csv': ('base_simulada/controle_erros/LOG_ERROS_INSERIDOS_FOLHA.csv', 'Log de controle dos erros inseridos na base de folha'),
}
# (arquivo, variavel, tipo, descricao, dominio / regra de geracao, real ou simulado)
D = [
 ('BASE', 'id_registro', 'texto', 'Identificador unico do registro no subrazao', 'IMB-0000001 a IMB-0532598 (sequencial apos embaralhamento)', 'SIMULADO'),
 ('BASE', 'entidade', 'texto', 'Entidade do grupo titular do item', 'PETRORECONCAVO (controladora) | SPE_TIETA (controlada 100%)', 'Estrutura REAL (NE 1); alocacao SIMULADA'),
 ('BASE', 'polo', 'texto', 'Polo/ativo de producao', 'REMANSO, REMANSO_BT_REC, MIRANGA, POTIGUAR, TIETA, MIDSTREAM_RN', 'Polos REAIS (NE 7.1 ii; FRE 1.1)'),
 ('BASE', 'campo', 'texto', 'Concessao (campo) ou local do item', '58 concessoes (nomes publicos quando disponiveis; "POT-CONC-xx (SINTETICO)" nos demais), blocos e UPGN Guamare', 'Parcialmente REAL'),
 ('BASE', 'uf', 'texto', 'Unidade da federacao', 'BA, SE, RN', 'REAL por polo'),
 ('BASE', 'poco_id', 'texto', 'Poco vinculado (classes desenvolvimento e ARO)', 'PRS-Cnn-nnnn; vazio nas demais classes', 'SIMULADO'),
 ('BASE', 'classe_ne7', 'texto', 'Classe do imobilizado conforme NE 7.1', '9 classes da NE 7.1 (maquinas, andamento, direito de producao, desenvolvimento, blocos, ARO, almoxarifado, adiantamentos, outros)', 'Classes REAIS'),
 ('BASE', 'conta_custo', 'texto', 'Conta analitica de custo (plano simulado)', 'ver PLANO_CONTAS_IMOBILIZADO_SIMULADO.csv', 'SIMULADO'),
 ('BASE', 'descricao_conta', 'texto', 'Descricao da conta de custo', 'texto do plano simulado', 'SIMULADO'),
 ('BASE', 'conta_depreciacao', 'texto', 'Conta redutora de depreciacao/deplecao', '1.2.03.9x.001; vazio para itens que nao depreciam', 'SIMULADO'),
 ('BASE', 'descricao_conta_depreciacao', 'texto', 'Descricao da conta redutora', 'texto', 'SIMULADO'),
 ('BASE', 'tipo_item', 'texto', 'Natureza do componente/item', 'ex.: PERFURACAO E COMPLETACAO; UNIDADE DE BOMBEIO MECANICO; TUBO DE REVESTIMENTO (LOTE)', 'SIMULADO'),
 ('BASE', 'origem', 'texto', 'Origem do registro', 'AQUISICAO_PETROBRAS, AQUISICAO_MAHA, AQUISICAO_BRAVA_JOA, RODADA_ANP, CAPEX_PROPRIO, COMPRA_MATERIAL, ARO_REMENSURACAO', 'Origens REAIS (NE 7, FRE 1.1); alocacao SIMULADA'),
 ('BASE', 'data_aquisicao', 'data', 'Data de aquisicao/capitalizacao/entrada em estoque', 'AAAA-MM-DD; posterior a aquisicao do polo', 'SIMULADO'),
 ('BASE', 'data_inicio_depreciacao', 'data', 'Inicio da depreciacao/deplecao', 'Linear: coerente com a acumulada (acumulada = custo x taxa x anos); UoP: = data de aquisicao', 'SIMULADO'),
 ('BASE', 'metodo_depreciacao', 'texto', 'Metodo de depreciacao', 'LINEAR | UOP_1P | UOP_PDP | NAO_DEPRECIA (NE 7.2)', 'Metodos REAIS'),
 ('BASE', 'taxa_anual_linear', 'decimal', 'Taxa anual (itens lineares)', 'Maquinas 10%; outros 4% a 20% (NE 7.2)', 'Taxas REAIS'),
 ('BASE', 'taxa_uop_campo_2025', 'decimal', 'Taxa de unidades produzidas do campo em 2025', 'producao 2025 / (reservas 31/12/2025 + producao 2025)', 'SIMULADO (a partir de totais reais)'),
 ('BASE', 'custo_historico', 'moeda (R$)', 'Custo historico do item', 'Rateio exato em centavos por entidade x classe; soma = NE 7.1 (arredondamento)', 'Totais REAIS; itens SIMULADOS'),
 ('BASE', 'depreciacao_acumulada', 'moeda (R$)', 'Depreciacao/deplecao acumulada (valor positivo da conta redutora)', 'Soma por entidade x classe = NE 7.1', 'Totais REAIS; itens SIMULADOS'),
 ('BASE', 'valor_contabil_liquido', 'moeda (R$)', 'Custo - depreciacao acumulada', 'soma consolidada = R$ 6.170.841 mil (CVM 1.02.03)', 'Total REAL; itens SIMULADOS'),
 ('BASE', 'depreciacao_2025', 'moeda (R$)', 'Depreciacao/deplecao reconhecida em 2025', 'UoP calibrada a NE 7.1; linear derivada das datas', 'SIMULADO'),
 ('BASE', 'status', 'texto', 'Situacao do item', 'EM_OPERACAO, EM_CONSTRUCAO, EM_ESTOQUE, ADIANTAMENTO, EXPLORACAO, BAIXADO', 'SIMULADO'),
 ('BASE', 'data_baixa', 'data', 'Data da baixa (se baixado)', 'vazio quando nao baixado', 'SIMULADO'),
 ('BASE', 'documento_origem', 'texto', 'Documento que suporta o registro', 'NF-/OI-/CTR-/CCV-/MEM-/ANP- + numero', 'SIMULADO'),
 ('BASE', 'fornecedor_id', 'texto', 'Fornecedor/contraparte', 'FORN-nnnnn; PETROBRAS; VENDEDORES_MAHA; ANP; 3R POTIGUAR (BRAVA); N/A (ESTIMATIVA)', 'SIMULADO (contrapartes reais nas aquisicoes)'),
 ('BASE', 'centro_custo', 'texto', 'Centro de custo', 'CC-<polo>-<classe>', 'SIMULADO'),
 ('BASE', 'garantia_judicial', 'texto', 'Bem dado em garantia judicial', 'S/N (uma sonda = S, NE 7.3)', 'Fato REAL; item SIMULADO'),
 ('BASE', 'usuario_cadastro', 'texto', 'Usuario que cadastrou (maker)', 'USR001-USR028', 'SIMULADO'),
 ('BASE', 'usuario_aprovacao', 'texto', 'Usuario que aprovou (checker)', 'USR001-USR028; deve ser diferente do maker', 'SIMULADO'),
 ('BASE', 'data_cadastro_sistema', 'data', 'Data de inclusao no sistema (trilha)', 'AAAA-MM-DD', 'SIMULADO'),
 ('BASE', 'data_ultima_alteracao', 'data', 'Data da ultima alteracao (trilha de auditoria)', 'AAAA-MM-DD', 'SIMULADO'),
 ('BASE', 'natureza_dado', 'texto', 'Marcador de dado sintetico', 'SINTETICO', 'Marcador'),
 ('INSPECAO', 'id_registro', 'texto', 'Chave para a base', 'igual a base', 'SIMULADO'),
 ('INSPECAO', 'documento_origem', 'texto', 'Documento inspecionado', 'igual a base', 'SIMULADO'),
 ('INSPECAO', 'documento_localizado', 'texto', 'Documento foi localizado?', 'S/N', 'SIMULADO'),
 ('INSPECAO', 'data_documento', 'data', 'Data constante do documento', 'AAAA-MM-DD', 'SIMULADO'),
 ('INSPECAO', 'valor_liquido_suportado', 'moeda (R$)', 'Valor liquido suportado pela evidencia (valor auditado no MUS)', '<= valor contabil', 'SIMULADO'),
 ('INSPECAO', 'observacao_inspecao', 'texto', 'Resultado da inspecao', 'texto', 'SIMULADO'),
 ('INSPECAO', 'natureza_dado', 'texto', 'Marcador', 'SINTETICO', 'Marcador'),
 ('CAMPOS', 'campo', 'texto', 'Nome da concessao', '58 concessoes', 'Parcialmente REAL'),
 ('CAMPOS', 'polo', 'texto', 'Polo', '5 polos', 'REAL'),
 ('CAMPOS', 'entidade', 'texto', 'Titular', 'PETRORECONCAVO | SPE_TIETA', 'REAL'),
 ('CAMPOS', 'ativo', 'texto', 'Ativo (agrupamento do RA)', 'BAHIA | POTIGUAR', 'REAL'),
 ('CAMPOS', 'uf', 'texto', 'UF', 'BA, SE, RN', 'REAL'),
 ('CAMPOS', 'data_aquisicao_polo', 'data', 'Data de aquisicao/inicio do polo', 'FRE 1.1', 'REAL (aproximada)'),
 ('CAMPOS', 'n_pocos', 'inteiro', 'Pocos simulados no campo', 'total 7.200', 'SIMULADO'),
 ('CAMPOS', 'producao_2025_boe', 'inteiro', 'Producao 2025 do campo', 'rateio de 13,4 mil boe/d (Bahia) e 13,1 mil boe/d (Potiguar) - RA', 'Total REAL; rateio SIMULADO'),
 ('CAMPOS', 'reservas_1P_31122025_boe', 'inteiro', 'Reservas provadas 1P', 'rateio de 87,3 MMboe (Bahia) e 56,1 MMboe (Potiguar) - RA/NSAI', 'Total REAL; rateio SIMULADO'),
 ('CAMPOS', 'reservas_PDP_31122025_boe', 'inteiro', 'Reservas PDP', 'rateio proporcional de 61,4 MMboe (RA)', 'Total REAL; rateio SIMULADO'),
 ('CAMPOS', 'taxa_uop_1P_2025', 'decimal', 'Taxa UoP 1P', 'producao/(reserva+producao)', 'SIMULADO'),
 ('CAMPOS', 'taxa_uop_PDP_2025', 'decimal', 'Taxa UoP PDP', 'producao/(reserva+producao)', 'SIMULADO'),
 ('CAMPOS', 'situacao_concessao', 'texto', 'Situacao', 'PRODUCAO', 'REAL (NE 17.1)'),
 ('CAMPOS', 'natureza_dado', 'texto', 'Marcador', 'texto', 'Marcador'),
 ('PLANO', 'conta_custo', 'texto', 'Codigo da conta', '1.2.03.xx.xxx', 'SIMULADO'), ('PLANO', 'descricao_conta', 'texto', 'Descricao', 'texto', 'SIMULADO'),
 ('PLANO', 'classe_ne7', 'texto', 'Classe da NE 7.1', '9 classes', 'REAL'), ('PLANO', 'conta_depreciacao', 'texto', 'Conta redutora', '1.2.03.9x.001', 'SIMULADO'),
 ('PLANO', 'natureza_dado', 'texto', 'Marcador', 'texto', 'Marcador'),
 ('POCOS', 'poco_id', 'texto', 'Identificador do poco', 'PRS-Cnn-nnnn', 'SIMULADO'), ('POCOS', 'campo', 'texto', 'Campo', '58 concessoes', 'Parcialmente REAL'),
 ('POCOS', 'polo', 'texto', 'Polo', '5 polos', 'REAL'), ('POCOS', 'entidade', 'texto', 'Titular', '2 entidades', 'REAL'), ('POCOS', 'uf', 'texto', 'UF', 'BA/SE/RN', 'REAL'),
 ('POCOS', 'data_entrada', 'data', 'Entrada do poco no grupo', 'aquisicao do polo ou perfuracao posterior', 'SIMULADO'),
 ('POCOS', 'origem_poco', 'texto', 'Origem', 'ADQUIRIDO_COM_O_POLO | PERFURADO_PELO_GRUPO', 'SIMULADO'), ('POCOS', 'natureza_dado', 'texto', 'Marcador', 'SINTETICO', 'Marcador'),
 ('LOG', 'seq', 'inteiro', 'Sequencial', '1..n', 'Controle'), ('LOG', 'tipo_erro', 'texto', 'Tipo do erro inserido', 'E01 a E11', 'Controle'),
 ('LOG', 'descricao', 'texto', 'Natureza do erro', 'texto', 'Controle'), ('LOG', 'campo_alterado', 'texto', 'Variavel alterada', 'texto', 'Controle'),
 ('LOG', 'valor_original', 'texto', 'Valor antes da insercao', 'texto', 'Controle'), ('LOG', 'valor_inserido', 'texto', 'Valor inserido', 'texto', 'Controle'),
 ('LOG', 'id_registro', 'texto', 'Registro afetado', 'IMB-...', 'Controle'), ('LOG', 'efeito_monetario_R$', 'moeda (R$)', 'Efeito no valor liquido (+ superavaliacao; - subavaliacao; 0 = nao monetario)', 'R$', 'Controle'),
 ('EMP', 'matricula', 'texto', 'Matricula do empregado', 'M00001...', 'SIMULADO'), ('EMP', 'entidade', 'texto', 'Empregador', '2 entidades', 'SIMULADO'),
 ('EMP', 'cpf_simulado', 'texto', 'CPF ficticio (nao e CPF real)', 'SIM-nnnnnnnnn', 'SIMULADO'), ('EMP', 'nivel', 'texto', 'Nivel hierarquico', 'LIDERANCA (249) | NAO_LIDERANCA (1.503) entre os ativos', 'Quantidades REAIS (FRE 10.1A)'),
 ('EMP', 'lotacao', 'texto', 'Lotacao', 'sede e polos', 'SIMULADO'), ('EMP', 'data_admissao', 'data', 'Admissao', 'AAAA-MM-DD', 'SIMULADO'),
 ('EMP', 'data_desligamento', 'data', 'Desligamento (vazio = ativo em 31/12/2025)', 'AAAA-MM-DD', 'SIMULADO'), ('EMP', 'salario_base_mensal', 'moeda (R$)', 'Salario base', 'calibrado a despesa de pessoal da NE 19', 'SIMULADO'),
 ('EMP', 'banco_agencia_conta', 'texto', 'Conta bancaria ficticia', 'bbb-aaaa-ccccccc', 'SIMULADO'), ('EMP', 'contrato_trabalho_arquivado', 'texto', 'Contrato arquivado', 'S/N', 'SIMULADO'),
 ('EMP', 'usuario_cadastro', 'texto', 'Maker (RH)', 'RH01-RH08', 'SIMULADO'), ('EMP', 'usuario_aprovacao', 'texto', 'Checker (RH)', 'RH01-RH08, diferente do maker', 'SIMULADO'),
 ('EMP', 'data_ultima_alteracao_salarial', 'data', 'Trilha: ultima alteracao salarial', 'AAAA-MM-DD', 'SIMULADO'), ('EMP', 'pct_ultimo_reajuste', 'decimal', 'Ultimo reajuste (%)', '0 a 65%', 'SIMULADO'),
 ('EMP', 'aprovador_reajuste', 'texto', 'Alcada que aprovou o reajuste', 'GERENCIA_RH | DIRETORIA_PESSOAS (acima de 30% exige diretoria)', 'SIMULADO'), ('EMP', 'natureza_dado', 'texto', 'Marcador', 'SINTETICO', 'Marcador'),
 ('FOL', 'id_lancamento', 'texto', 'Id do lancamento', 'FOL-nnnnnnn', 'SIMULADO'), ('FOL', 'competencia', 'texto', 'Competencia', '2025-01 a 2025-12; 2025-13 = 13o', 'SIMULADO'),
 ('FOL', 'matricula', 'texto', 'Matricula paga', 'chave do cadastro', 'SIMULADO'), ('FOL', 'rubrica', 'texto', 'Rubrica', 'SALARIO, ADICIONAL_PERICULOSIDADE, HORAS_EXTRAS, 13O_SALARIO, FERIAS_+1/3, PLR_2024_PAGA, VERBAS_RESCISORIAS', 'SIMULADO'),
 ('FOL', 'valor_bruto', 'moeda (R$)', 'Valor bruto', 'R$', 'SIMULADO'), ('FOL', 'data_pagamento', 'data', 'Data do pagamento', 'AAAA-MM-DD', 'SIMULADO'),
 ('FOL', 'conta_credito', 'texto', 'Conta creditada', 'conta do cadastro', 'SIMULADO'), ('FOL', 'encargos_patronais', 'moeda (R$)', 'INSS patronal + FGTS (36,8%)', 'R$', 'SIMULADO'),
 ('FOL', 'beneficios', 'moeda (R$)', 'Beneficios (11,5% da folha bruta)', 'R$', 'SIMULADO'), ('FOL', 'natureza_dado', 'texto', 'Marcador', 'SINTETICO', 'Marcador'),
 ('OBR', 'id_obrigacao', 'texto', 'Id da obrigacao', 'PROV-nnnnnn', 'SIMULADO'), ('OBR', 'matricula', 'texto', 'Empregado', 'chave do cadastro', 'SIMULADO'),
 ('OBR', 'tipo_obrigacao', 'texto', 'Natureza', 'salario de dezembro, encargos, ferias, PLR 2025, outras, bonus de retencao', 'SIMULADO'),
 ('OBR', 'valor', 'moeda (R$)', 'Saldo em 31/12/2025', 'soma = R$ 88.753 mil (CVM 2.01.01)', 'Total REAL; itens SIMULADOS'),
 ('OBR', 'data_vencimento', 'data', 'Vencimento', 'ate 31/12/2026 = circulante; a partir de 01/01/2027 = nao circulante', 'SIMULADO'),
 ('OBR', 'classificacao_registrada', 'texto', 'Classificacao no balanco', 'CIRCULANTE', 'SIMULADO'), ('OBR', 'natureza_dado', 'texto', 'Marcador', 'SINTETICO', 'Marcador'),
 ('LOGF', 'tipo_erro', 'texto', 'Tipo', 'F01 a F06', 'Controle'), ('LOGF', 'chave', 'texto', 'Matricula ou id_obrigacao afetado', 'texto', 'Controle'),
 ('LOGF', 'descricao', 'texto', 'Natureza do erro', 'texto', 'Controle'), ('LOGF', 'efeito_monetario_R$', 'moeda (R$)', 'Efeito monetario', 'R$', 'Controle'),
]
SIG = {'BASE': 'BASE_IMOBILIZADO_RECV3_SINTETICA.csv.gz', 'INSPECAO': 'INSPECAO_DOCUMENTAL_SIMULADA.csv.gz', 'CAMPOS': 'PARAMETROS_CAMPOS_SIMULADOS.csv',
       'PLANO': 'PLANO_CONTAS_IMOBILIZADO_SIMULADO.csv', 'POCOS': 'CADASTRO_POCOS_SIMULADO.csv', 'LOG': 'LOG_ERROS_INSERIDOS.csv', 'EMP': 'CADASTRO_EMPREGADOS_SIMULADO.csv',
       'FOL': 'FOLHA_PAGAMENTOS_2025_SIMULADA.csv', 'OBR': 'OBRIGACOES_TRABALHISTAS_31122025_SIMULADAS.csv', 'LOGF': 'LOG_ERROS_INSERIDOS_FOLHA.csv'}
dic = pd.DataFrame(D, columns=['arquivo', 'variavel', 'tipo', 'descricao', 'dominio_regra', 'real_ou_simulado'])
dic['arquivo'] = dic.arquivo.map(SIG)
# conferencia: variaveis documentadas x colunas reais dos arquivos
prob = []; linhas = {}
for nome, (path, desc) in ARQ.items():
    df = pd.read_csv(path, sep=';', nrows=5, encoding='utf-8-sig' if not path.endswith('.gz') else 'utf-8')
    doc = set(dic[dic.arquivo == nome].variavel); real = set(df.columns)
    if doc != real: prob.append((nome, sorted(real - doc), sorted(doc - real)))
    with (gzip.open(path, 'rt', encoding='utf-8') if path.endswith('.gz') else open(path, encoding='utf-8-sig')) as f:
        linhas[nome] = sum(1 for _ in f) - 1
assert not prob, prob
wb = Workbook(); ws = wb.active; ws.title = 'Arquivos'
ws.append(['Arquivo', 'Caminho', 'Conteudo', 'Linhas']); [ws.append([k, v[0], v[1], linhas[k]]) for k, v in ARQ.items()]
ws2 = wb.create_sheet('Variaveis'); ws2.append(list(dic.columns)); [ws2.append(list(r)) for r in dic.itertuples(index=False)]
for w, larg in ((ws, [44, 62, 90, 10]), (ws2, [44, 30, 12, 60, 80, 34])):
    for j, x in enumerate(larg, 1): w.column_dimensions[get_column_letter(j)].width = x
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
    for row in w.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save('base_simulada/DICIONARIO_VARIAVEIS.xlsx')
md = ['# Dicionario de variaveis - bases simuladas (DADOS SINTETICOS)', '', 'Preparador: matricula 211059370. Gerado por `codigo/gerar_dicionario.py`, que tambem confere se cada variavel documentada existe no arquivo.', '',
      '## Arquivos', '', '| Arquivo | Linhas | Conteudo |', '|---|---:|---|'] + [f'| `{v[0]}` | {linhas[k]:,} | {v[1]} |'.replace(',', '.') for k, v in ARQ.items()]
for nome in ARQ:
    md += ['', f'## {nome}', '', '| Variavel | Tipo | Descricao | Dominio / regra | Real ou simulado |', '|---|---|---|---|---|']
    md += [f'| `{r.variavel}` | {r.tipo} | {r.descricao} | {r.dominio_regra} | {r.real_ou_simulado} |' for r in dic[dic.arquivo == nome].itertuples()]
open('base_simulada/DICIONARIO_VARIAVEIS.md', 'w', encoding='utf-8').write('\n'.join(md) + '\n')
print('ok dicionario', {k: linhas[k] for k in linhas})
