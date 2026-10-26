# -*- coding: utf-8 -*-
"""
Etapa 2 - Diagnostico da companhia (insumo PT 315, 320, 520, 540, 550, 570, 600)
Preparador: matricula 211059370 | PetroReconcavo S.A. (RECV3) | exercicio 2025
Gera planilhas/DIAG_RECV3.xlsx. Valores em R$ mil.
Fontes: F01 (DFP 2025 em PDF) e F03/F04 (dados abertos CVM) - ver planilhas/REG_FONTES_RECV3.xlsx.
As celulas de calculo usam FORMULAS do Excel para permitir conferencia e reproducao.
Uso: python codigo/etapa2_diagnostico.py   (executar na raiz P2_PetroReconcavo)
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CVM = 'fontes/cvm_dados_abertos/RECV3_dfp_cia_aberta_{}_{}.csv'
HEAD = PatternFill('solid', fgColor='1F4E78'); SUB = PatternFill('solid', fgColor='DDEBF7'); ALERT = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);-'; PCT = '0.0%'
RODAPE = 'Preparador: matricula 211059370 | Data: 22/10/2026 | Valores em R$ mil | Revisao: ______ (campo reservado; nao preenchido)'


def ler(dem, ano):
    """lista de (codigo, descricao, valor exercicio, valor comparativo) da DFP do ano"""
    d = {}
    with open(CVM.format(dem, ano), encoding='utf-8') as f:
        for r in csv.DictReader(f, delimiter=';'):
            if 'CD_CONTA' not in r:
                continue
            k = r['CD_CONTA']
            d.setdefault(k, [r['DS_CONTA'], None, None])
            v = float(r['VL_CONTA'])
            if r['ORDEM_EXERC'] == 'ÚLTIMO':
                d[k][1] = v
            else:
                d[k][2] = v
    ordem = lambda k: [int(p) for p in k.split('.')]
    return [(k, d[k][0], d[k][1], d[k][2]) for k in sorted(d, key=ordem)]


def cab(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font = Font(bold=True, color='FFFFFF'); x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BORDER
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def titulo(ws, txt, sub=None):
    ws['A1'] = txt; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = sub or RODAPE; ws['A2'].font = Font(italic=True, size=9)


def fmt(ws, r1, r2, c1, c2, nf=NUM):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(row=r, column=c); x.number_format = nf; x.border = BORDER


def wrap(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(row=r, column=c)
            x.alignment = Alignment(wrap_text=True, vertical='top'); x.border = BORDER


wb = Workbook(); wb.remove(wb.active)

# ===================================================== 1. Ranking de materialidade do ativo (escolha da conta)
ws = wb.create_sheet('Ranking_Ativo')
titulo(ws, 'Ranking das rubricas do ativo consolidado em 31/12/2025 - escolha da conta da base simulada (item 3 do enunciado)')
cab(ws, 4, ['#', 'Rubrica do balanco (F01 p.51)', 'Nota', '31/12/2025', '% ativo total', '31/12/2024', '% ativo total 2024', 'Multiplo s/ 2a maior'],
    [5, 52, 8, 16, 13, 16, 13, 14])
ativo = [('Imobilizado e intangivel', '7', 6254225, 5561314), ('Aplicacoes financeiras (CP)', '4', 1395510, 761939),
         ('Contas a receber de clientes (CP)', '5', 253967, 361095), ('Caixa e equivalentes de caixa', '4', 229508, 295548),
         ('Tributos a recuperar (CP)', '-', 146817, 96616), ('Outros ativos (CP)', '-', 92524, 43886), ('Outros ativos (LP)', '-', 81912, 46540),
         ('Tributos a recuperar (LP)', '-', 69109, 66820), ('Contas a receber de clientes (LP)', '5', 65231, 58145),
         ('Instrumentos financeiros derivativos (CP)', '14', 33771, 575), ('Direito de uso em arrendamento', '-', 24883, 22338),
         ('Tributos diferidos', '10', 8113, 97025), ('Estoques', '-', 6139, 9766), ('Aplicacoes financeiras (LP)', '4', 5022, 15964),
         ('Instrumentos financeiros derivativos (LP)', '14', 1284, 0)]
ativo.sort(key=lambda x: -x[2])
r = 5
for i, (n, ne, v1, v0) in enumerate(ativo, 1):
    ws.append([i, n, ne, v1, f'=D{r}/D$21', v0, f'=F{r}/F$21', f'=D{r}/D$6'])
    if i == 1:
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = SUB; ws.cell(row=r, column=c).font = Font(bold=True)
    r += 1
ws.append(['', 'Soma das rubricas', '', '=SUM(D5:D19)', '=D20/D21', '=SUM(F5:F19)', '=F20/F21', ''])
ws.append(['', 'TOTAL DO ATIVO (F01 p.51; CVM conta 1)', '', 8668015, '', 7437571, '', ''])
ws.append(['', 'Checagem (soma - total) - deve ser zero', '', '=D20-D21', '', '=F20-F21', '', ''])
fmt(ws, 5, 22, 4, 4); fmt(ws, 5, 22, 6, 6); fmt(ws, 5, 22, 5, 5, PCT); fmt(ws, 5, 22, 7, 7, PCT); fmt(ws, 5, 19, 8, 8, '0.0"x"')
for i, t in enumerate([
        'Conclusao: Imobilizado e intangivel (NE 7) e a rubrica de maior materialidade do ativo: 72,2% do ativo total consolidado (74,8% em 2024),',
        'cerca de 4,5 vezes a segunda maior rubrica (aplicacoes financeiras). O intangivel (software) e residual (R$ 83.384; CVM 1.02.04):',
        'o imobilizado (CVM 1.02.03 = R$ 6.170.841) responde por 98,7% da rubrica. A base simulada sera o subrazao analitico do imobilizado (aba NE7_Imobilizado).',
        'Mesma conclusao na controladora: imobilizado e intangivel R$ 5.689.089 = 66,2% do ativo total de R$ 8.597.497.']):
    ws.cell(row=24 + i, column=2, value=t)

# ===================================================== 2. NE 7 - Imobilizado consolidado
ws = wb.create_sheet('NE7_Imobilizado')
titulo(ws, 'NE 7.1 - Imobilizado e intangivel consolidado: composicao e movimentacao 2025 (F01 p.69) - transcricao + recalculo')
cab(ws, 4, ['Classe', 'Custo 31/12/2024', 'Adicoes', 'Baixas', 'Transferencias', 'UPGN Guamare', 'Custo 31/12/2025 (NE)', 'Custo final recalculado',
            'Dif.', 'Deprec./deplecao acum. 31/12/2025 (NE)', 'Valor liquido 31/12/2025', '% do imobilizado liquido', 'Metodo (NE 7.2)'],
    [36, 14, 12, 12, 13, 13, 14, 14, 8, 16, 15, 12, 14])
imob = [('Maquinas e equipamentos', 309092, 1989, 0, 122441, 284999, 718521, -103788, '10% a.a.'),
        ('UPGN Guamare (conta transitoria)', 0, 312790, 0, 36849, -349639, 0, 0, '-'),
        ('Imobilizados em andamento', 141241, 88691, -9212, -158390, 9841, 72171, 0, 'nao deprecia'),
        ('Direito de producao de oleo e gas', 2973528, 0, -42199, 0, 0, 2931329, -922414, 'UoP (reservas 1P)'),
        ('Desenvolvimento de campos', 4180242, 1006863, -5545, 5116, 0, 5186676, -2422704, 'UoP (reservas PDP)'),
        ('Blocos exploratorios', 20037, 120, 0, 0, 0, 20157, 0, 'nao depleta'),
        ('Abandono de poco (ativo ARO)', 79091, 8340, -8128, 0, 0, 79303, -41848, 'UoP'),
        ('Almoxarifado para inversoes fixas', 502638, 223654, -166296, -82345, 0, 477651, 0, 'nao deprecia'),
        ('Adiantamentos', 46219, 32906, -8051, -52655, 0, 18419, 0, '-'),
        ('Outros', 105066, 735, -300, 47594, 46230, 199325, -41957, '4% a 25% a.a.')]
r = 5
for n, c0, ad, bx, tr, up, c1, dep, met in imob:
    ws.append([n, c0, ad, bx, tr, up, c1, f'=SUM(B{r}:F{r})', f'=G{r}-H{r}', dep, f'=G{r}+J{r}', f'=K{r}/K$15', met]); r += 1
ws.append(['Total imobilizado', '=SUM(B5:B14)', '=SUM(C5:C14)', '=SUM(D5:D14)', '=SUM(E5:E14)', '=SUM(F5:F14)', '=SUM(G5:G14)',
           '=SUM(H5:H14)', '=G15-H15', '=SUM(J5:J14)', '=SUM(K5:K14)', '=K15/K15', ''])
ws.append(['Imobilizado liquido - CVM conta 1.02.03 (F03)', '', '', '', '', '', 9703552, '', '', -3532711, 6170841, '', ''])
ws.append(['Diferenca (calculado - fonte)', '', '', '', '', '', '=G15-G16', '', '', '=J15-J16', '=K15-K16', '', ''])
ws.append(['Intangivel - software', 32955, 3639, 0, 67878, 8570, 113042, '=SUM(B18:F18)', '=G18-H18', -29658, '=G18+J18', '', '20% a.a.'])
ws.append(['Total imobilizado e intangivel (calculado)', '', '', '', '', '', '', '', '', '', '=K15+K18', '', ''])
ws.append(['Saldo no balanco (F01 p.51)', '', '', '', '', '', '', '', '', '', 6254225, '', ''])
ws.append(['Diferenca', '', '', '', '', '', '', '', '', '', '=K19-K20', '', ''])
fmt(ws, 5, 21, 2, 11); fmt(ws, 5, 15, 12, 12, PCT)
for c in range(1, 14):
    ws.cell(row=15, column=c).font = Font(bold=True)
for i, t in enumerate([
        'Custo total 31/12/2025 na NE 7.1 consolidada = 9.703.552 e depreciacao acumulada = (3.532.711) (F01 p.69).',
        'D&A 2025 (NE 7.1 consolidada): imobilizado (671.868) + intangivel (16.802) = (688.670); NE 19/DFC: D&A total 721.077 (inclui direito de uso) - conciliar no PT 520.',
        'Custo de aquisicao por polo (NE 7.1 ii, p.70): Remanso 95.629; Remanso BT-REC 1.248; Miranga 1.247.506; Potiguar 1.507.572; Tieta 79.374 = 2.931.329.',
        'Transferencias nao zeram (custo -81.390; deprec. -62.640): a nota (a) atribui a transferencias de creditos tributarios e itens de alto giro, sem quantificar - investigar (PT 500/520).',
        'Farm-out Mandacaru (NE 7.4, p.71): 50% de 7 concessoes por US$ 5 milhoes; politica: manter os ativos no imobilizado - julgamento contabil relevante.',
        'Impairment (NE 3.4, p.62): administracao nao identificou indicativos, embora valor de mercado (R$ 3,3 bi; RA p.28) < PL (R$ 4,33 bi) - ver aba Achados_Preliminares.',
        'Garantia (NE 7.3, p.70): uma sonda de perfuracao dada em garantia de execucao fiscal.']):
    ws.cell(row=23 + i, column=1, value=t)

# ===================================================== 3. Demonstracoes (CVM) com AV/AH
DEMS = [('BPA', 'Balanco - Ativo'), ('BPP', 'Balanco - Passivo e PL'), ('DRE', 'Resultado'), ('DRA', 'Resultado abrangente'),
        ('DFC_MI', 'Fluxo de caixa (metodo indireto)'), ('DVA', 'Valor adicionado')]
BASE = {'BPA': '1', 'BPP': '2', 'DRE': '3.01', 'DVA': '7.08'}
for esc in ('con', 'ind'):
    for dem, nome in DEMS:
        ws = wb.create_sheet(f'{dem}_{esc}')
        titulo(ws, f'{nome} - {"Consolidado" if esc == "con" else "Controladora"} - 2025 x 2024',
               'Fonte: dados abertos CVM (F03), DFP 31/12/2025 v1, conciliada ao PDF F01 p.51-56 | ' + RODAPE)
        cab(ws, 4, ['Codigo CVM', 'Conta', '31/12/2025', '31/12/2024', 'AV 2025', 'AV 2024', 'AH (var. R$)', 'AH (var. %)'], [14, 64, 16, 16, 9, 9, 15, 10])
        linhas = [l for l in ler(f'{dem}_{esc}', 2025) if not (abs(l[2] or 0) < 0.5 and abs(l[3] or 0) < 0.5)]
        r = 5; base_row = None
        for cod, ds, v1, v0 in linhas:
            ws.cell(row=r, column=1, value=cod); ws.cell(row=r, column=2, value=('   ' * cod.count('.')) + ds)
            ws.cell(row=r, column=3, value=v1); ws.cell(row=r, column=4, value=v0)
            if cod == BASE.get(dem):
                base_row = r
            if cod.count('.') <= 1:
                for c in range(1, 9):
                    ws.cell(row=r, column=c).font = Font(bold=True)
            r += 1
        last = r - 1
        for i in range(5, last + 1):
            if base_row:
                ws.cell(row=i, column=5, value=f'=IFERROR(C{i}/C${base_row},"")'); ws.cell(row=i, column=6, value=f'=IFERROR(D{i}/D${base_row},"")')
            ws.cell(row=i, column=7, value=f'=C{i}-D{i}'); ws.cell(row=i, column=8, value=f'=IFERROR(C{i}/D{i}-1,"")')
        fmt(ws, 5, last, 3, 4); fmt(ws, 5, last, 7, 7); fmt(ws, 5, last, 5, 6, PCT); fmt(ws, 5, last, 8, 8, PCT)
        ws.freeze_panes = 'C5'
        ws.cell(row=last + 2, column=1, value='AV: base = ativo total / passivo total / receita liquida / valor adicionado a distribuir. AH: 2025 versus 2024.').font = Font(italic=True, size=9)

# ===================================================== 4. Consistencia dos comparativos (PT 510)
ws = wb.create_sheet('Comparativos_2024')
titulo(ws, 'Consistencia: comparativo 2024 apresentado na DFP 2025 x saldo final da DFP 2024 (insumo PT 510 - saldos iniciais)')
cab(ws, 4, ['Demonstracao', 'Codigo', 'Conta', '2024 na DFP 2025 (F03)', '2024 na DFP 2024 v2 (F04)', 'Diferenca'], [12, 14, 58, 20, 20, 14])
r = 5; ndif = 0
for dem in ('BPA_con', 'BPP_con', 'DRE_con', 'DFC_MI_con', 'BPA_ind', 'BPP_ind', 'DRE_ind'):
    a = {c: (d, v0) for c, d, v1, v0 in ler(dem, 2025)}
    b = {c: (d, v1) for c, d, v1, v0 in ler(dem, 2024)}
    for c in sorted(set(a) | set(b), key=lambda x: [int(p) for p in x.split('.')]):
        if c.count('.') > 2:
            continue
        va = (a.get(c, ('', 0))[1] or 0); vb = (b.get(c, ('', 0))[1] or 0)
        if abs(va) < .5 and abs(vb) < .5:
            continue
        ws.append([dem, c, (a.get(c) or b.get(c))[0], va, vb, f'=D{r}-E{r}'])
        if abs(va - vb) > .5:
            ndif += 1
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = ALERT
        r += 1
fmt(ws, 5, r - 1, 4, 6)
ws.cell(row=r + 1, column=1, value=f'Linhas com diferenca (> R$ 0,5 mil): {ndif}. Diferencas indicam reclassificacao/reapresentacao do comparativo a ser explicada no PT 510.').font = Font(bold=True)
NDIF = ndif

# ===================================================== 5. Debentures e covenants
ws = wb.create_sheet('NE9_Debentures')
titulo(ws, 'NE 9 - Debentures (F01 p.71-75) e recalculo independente dos covenants (insumo PT 570)')
cab(ws, 4, ['Emissao / serie', '31/12/2025', '31/12/2024', 'Indexador / taxa (NE 9)', 'Vencimento', 'Liquidacao'], [44, 15, 15, 38, 14, 26])
deb = [('1a emissao - serie 1', 811874, 777481, 'IPCA + 7,32% (Lei 12.431)', '15/05/2031', '04/06/2024'),
       ('1a emissao - serie 2', 381973, 381789, '12,88% a.a. (pre)', '15/05/2031', '04/06/2024'),
       ('1a emissao - custos a amortizar', -25185, -29724, '', '', ''),
       ('2a emissao', 670449, 664190, 'nao informado na NE (swap: 6,16% em US$)', '15/10/2029', '11/10/2024'),
       ('2a emissao - custos a amortizar', -1083, -1415, '', '', ''),
       ('3a emissao', 539217, 0, 'nao informado na NE (swap: 5,66% em US$)', '04/07/2032', '04/07/2025'),
       ('3a emissao - custos a amortizar', -2460, 0, '', '', ''),
       ('4a emissao - serie 1 (rotulo da tabela NE 9.1)', 225208, 0, 'IPCA + 7,40% (NE 9) / 7,4564% (NE 25)', '15/12/2032', '18/12/2025 (NE 9) / 29/12/2025 (NE 25)'),
       ('4a emissao - serie 2 (rotulo da tabela NE 9.1)', 525492, 0, 'IPCA + 7,40% (NE 9) / 7,2833% (NE 25)', '15/12/2035', 'idem'),
       ('4a emissao - custos a amortizar', -20469, 0, '', '', '')]
for d in deb:
    ws.append(list(d))
ws.append(['Total (recalculado)', '=SUM(B5:B14)', '=SUM(C5:C14)', '', '', ''])
ws.append(['Total divulgado (NE 9.1)', 3105016, 1792321, '', '', ''])
ws.append(['Diferenca', '=B15-B16', '=C15-C16', '', '', ''])
fmt(ws, 5, 17, 2, 3)
ws['A19'] = 'Covenants financeiros (NE 9.3, p.75) - recalculo com dados publicos'; ws['A19'].font = Font(bold=True)
cab(ws, 20, ['Componente', 'Valor', 'Fonte / formula', 'Limite', 'Status', ''])
cov = [('Caixa e equivalentes', 229508, 'F01 BP / NE 4.1'), ('Aplicacoes financeiras CP + LP', 1400532, 'F01 NE 4.2'),
       ('Caixa livre (covenant: >= R$ 100.000 mil)', '=B21+B22', 'soma'), ('Divida bruta - debentures', 3105016, 'F01 NE 9.1'),
       ('Derivativos passivos (LP)', 88449, 'F01 BP'), ('Derivativos ativos (CP + LP)', -35055, 'F01 BP (33.771 + 1.284)'),
       ('Divida liquida (calculo do auditor, sem arrendamentos)', '=B24+B25+B26-B23', 'debentures + derivativos liquidos - caixa livre'),
       ('Lucro operacional consolidado', 721579, 'F01 DRE p.52'), ('Depreciacao, amortizacao e deplecao', 721077, 'F01 NE 19 / DFC'),
       ('EBITDA (calculo do auditor)', '=B28+B29', 'lucro operacional + D&A'),
       ('Divida liquida / EBITDA (covenant: <= 3,00x)', '=B27/B30', 'calculado'),
       ('Divida liquida / EBITDA divulgada pela administracao', 1.10, 'F01 RA p.5 (divida liquida R$ 1,6 bi)'),
       ('PV-10 reservas 1P (US$ mil)', 1900000, 'F01 RA 3.3 / F14 NSAI'), ('PTAX 31/12/2025 (R$/US$)', 5.5024, 'F01 NE 21 nota (a) p.98'),
       ('PV-10 1P em R$ mil', '=B33*B34', 'conversao'), ('PV-10 / divida bruta (covenant: >= 1,50x)', '=B35/B24', 'calculado')]
for n, v, f in cov:
    ws.append([n, v, f, '', '', ''])
ws['D23'] = 100000; ws['E23'] = '=IF(B23>=D23,"cumpre","NAO cumpre")'
ws['D31'] = 3; ws['E31'] = '=IF(B31<=D31,"cumpre","NAO cumpre")'
ws['D36'] = 1.5; ws['E36'] = '=IF(B36>=D36,"cumpre","NAO cumpre")'
fmt(ws, 21, 36, 2, 2)
for c in ('B31', 'B32', 'B36', 'D31', 'D36'):
    ws[c].number_format = '0.00"x"'
ws['B34'].number_format = '0.0000'
ws['A38'] = 'Limitacao: escrituras nao disponiveis no RI; definicoes contratuais de Divida Liquida e EBITDA nao verificadas (F31-F34 a examinar no PT 570).'
ws['A39'] = 'O calculo do auditor difere do divulgado (R$ 1,6 bi) por itens nao identificados (provavelmente arrendamentos e valores a pagar por aquisicoes).'
ws['A40'] = 'Cronograma do nao circulante (NE 9.2): 2028 308.537; 2029 710.737; 2030 553.048; 2031 817.695; 2032+ 650.085 = 3.040.102. Sem principal relevante em 2026-2027.'

# ===================================================== 6. Mapa de estimativas
ws = wb.create_sheet('Estimativas')
titulo(ws, 'Mapa das estimativas contabeis (NE 3) - insumo PT 540 e PT 315')
cab(ws, 4, ['Estimativa', 'Saldo / efeito 2025 (R$ mil)', 'Nota / pagina F01', 'Premissas-chave divulgadas', 'Grau de incerteza / julgamento', 'Candidata ao PT 540?'],
    [30, 26, 18, 60, 40, 16])
est = [('Provisao para abandono de pocos (ARO)', 'Passivo 145.192 (CP 4.728; LP 140.464); ativo ARO liquido 37.455', 'NE 3.3 p.61; NE 13 p.82-83',
        'Taxa nominal 10,56% a.a. (2024: 10,92%); inflacao 3,56% a.a. (2024: 3,66%); epoca do abandono pela projecao de producao; custo por servico',
        'Alto: longo prazo, regulacao subjetiva (ANP); unico PAA do auditor', 'SIM (principal)'),
       ('Reservas -> deplecao por unidades produzidas', 'Deplecao 2025: direito de producao 132.263 + desenvolvimento 485.386 + ARO 1.174 = 618.823',
        'NE 3.2 p.61; NE 7 p.69', 'Reservas 1P para custo de aquisicao; PDP para desenvolvimento; NSAI (F14): 1P 143,3 MMboe; PDP 61,4 MMboe',
        'Alto: depende de especialista (NSAI) e do preco do Brent', 'SIM (alternativa)'),
       ('Impairment do imobilizado (6 UGCs)', 'Sem perda reconhecida; imobilizado liquido 6.170.841', 'NE 3.4 p.62',
        'Administracao nao identificou indicativos; UGC = grupo de campos', 'Alto: indicativo externo (valor de mercado < PL; Brent -14%)', 'SIM (risco significativo)'),
       ('IR/CS diferidos', 'Liquido passivo (52.211); ativo bruto 270.697', 'NE 3.1; NE 10.2 p.77',
        'Realizacao conforme vencimento de derivativos, abandono e provisoes; SUDENE 75% ate 2032; deplecao acelerada (Lei 13.586/2017)', 'Medio', 'Nao'),
       ('Valor justo de derivativos (swaps, NDF, collar)', 'Ativo 35.055; passivo 88.449; resultado +366.796', 'NE 3.5; NE 14; NE 20 p.94', 'Curvas de mercado',
        'Medio, mas alto impacto no resultado', 'Nao (testar no PT 330)'),
       ('Pagamento baseado em acoes', 'Programas ILP 2022-2024', 'NE 3.6; NE 15.7', 'Valor justo na outorga; parcela TSR', 'Baixo', 'Nao'),
       ('Creditos em disputa (gas 2022) / perdas de credito', 'Outros a receber liquidos 55.917 (LP)', 'NE 5.1 (i) p.64; NE 12.3 p.82',
        'Melhor estimativa de realizacao; arbitragem CCI suspensa para mediacao', 'Alto qualitativo (litigio confidencial)', 'Nao (PT 501)'),
       ('Provisoes para processos judiciais', 'Provavel 47.946; possivel 181.206', 'NE 12 p.80-81', 'Opiniao de assessores juridicos', 'Medio', 'Nao (PT 501)'),
       ('Contraprestacao contingente (earnouts)', 'Zerada em 2025 (2024: 27.308)', 'NE 11 p.79', 'Curva futura do Brent', 'Baixo em 2025', 'Nao')]
for e in est:
    ws.append(list(e))
wrap(ws, 5, 4 + len(est), 1, 6)

# ===================================================== 7. Litigios, partes relacionadas e receita
ws = wb.create_sheet('Litigios_PR_Receita')
titulo(ws, 'Litigios (NE 12), partes relacionadas (NE 16) e receita (NE 18)')
for j, w in enumerate([48, 18, 18, 18, 18], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
r = 4
ws.cell(row=r, column=1, value='NE 12 - Processos judiciais - consolidado (F01 p.80-81)').font = Font(bold=True); r += 1
cab(ws, r, ['Natureza', 'Provavel 2025', 'Provavel 2024', 'Possivel 2025', 'Possivel 2024']); r += 1; r0 = r
for row in [('Trabalhistas', 3214, 4810, 6283, 4098), ('Fiscais', 1698, 1472, 82797, 57300),
            ('Regulatorios (ANP - PEM nao executado, SPE Tieta)', 43034, 41641, 82965, 49891), ('Civeis', 0, 0, 9161, 8360)]:
    ws.append(list(row)); r += 1
ws.append(['Total (recalculado)'] + [f'=SUM({get_column_letter(c)}{r0}:{get_column_letter(c)}{r - 1})' for c in range(2, 6)])
ws.append(['Total divulgado', 47946, 47923, 181206, 119649]); r += 2
fmt(ws, r0, r - 1, 2, 5); r += 1
ws.cell(row=r, column=1, value='NE 16 - Partes relacionadas (F01 p.90-91)').font = Font(bold=True); r += 1
cab(ws, r, ['Parte / natureza', 'Saldo 2025 (controladora)', 'Saldo 2024 (controladora)', 'Resultado 2025 (controladora)', 'Resultado 2025 (consolidado)']); r += 1; r0 = r
for row in [('SPE Tieta - outros ativos (intragrupo)', 3165, 27399, 31667, 0), ('SPE Tieta - fornecedores (intragrupo)', 2597, 3314, 0, 0),
            ('Grupo PERBRAS (acionista) - sondas e servicos', 5027, 2054, -4503, -4509), ('Grupo PetroSantander (acionista) - assistencia tecnica', 0, 1, -961, -961),
            ('Rateios de gastos corporativos (intragrupo)', 0, 0, 44763, 0), ('Remuneracao da administracao (total, c/ encargos)', 0, 0, -30694, -30694)]:
    ws.append(list(row)); r += 1
fmt(ws, r0, r - 1, 2, 5); r += 1
ws.cell(row=r, column=1, value='NE 18 - Receita - consolidado (F01 p.93)').font = Font(bold=True); r += 1
cab(ws, r, ['Componente', '2025', '2024', 'AH %', '% receita bruta 2025']); r += 1; r0 = r
for row in [('Venda de petroleo', 2113133, 2440303), ('Venda de gas e subprodutos', 1554086, 1480337), ('Prestacao de servicos', 41583, 43315),
            ('Contrato de hedge', 882, -127212)]:
    ws.append(list(row) + [f'=IFERROR(B{r}/C{r}-1,"")', f'=B{r}/B{r0 + 4}']); r += 1
ws.append(['Receita bruta', f'=SUM(B{r0}:B{r - 1})', f'=SUM(C{r0}:C{r - 1})', f'=B{r}/C{r}-1', '']); r += 1
ws.append(['(-) Deducoes', -552075, -572189, f'=B{r}/C{r}-1', '']); r += 1
ws.append(['Receita liquida (recalculada)', f'=B{r - 2}+B{r - 1}', f'=C{r - 2}+C{r - 1}', f'=B{r}/C{r}-1', '']); r += 1
ws.append(['Receita liquida divulgada (DRE)', 3157609, 3264554, '', '']); r += 1
fmt(ws, r0, r - 1, 2, 3); fmt(ws, r0, r - 1, 4, 5, PCT); r += 1
for t in ['Concentracao de clientes (NE 21, p.97): ~83% da receita com clientes que representam > 10% cada; os tres maiores = 22%, 28% e 33% da receita.',
          'Exposicao cambial (NE 21, p.97): 97% da receita bruta vinculada ao US$ (Brent ou contratos em US$).',
          'Royalties (NE 17.1, p.92): R$ 198.434 (2024: 196.246) + proprietarios de terra R$ 29.260. Royalties no custo (NE 19): R$ 227.694.',
          'Arbitragem CCI sobre contratos de venda de gas (NE 12.3, p.82): confidencial, suspensa para mediacao; ativos/passivos relacionados podem variar.',
          'Dividendos de R$ 300.000 aprovados em 18/12/2025 (NE 15.6, p.88) com base na Lei 15.270/2025: R$ 100.000 CP e R$ 200.000 LP (F01 BP).']:
    ws.cell(row=r, column=1, value=t); r += 1

# ===================================================== 8. Indicadores
ws = wb.create_sheet('Indicadores')
titulo(ws, 'Indicadores de liquidez, endividamento, rentabilidade e caixa - consolidado (insumo PT 520 e PT 570)')
cab(ws, 4, ['Item', '2025', '2024', 'Fonte'], [60, 16, 16, 46])
p25 = {c: v1 for c, d, v1, v0 in ler('BPP_con', 2025)}; p24 = {c: v0 for c, d, v1, v0 in ler('BPP_con', 2025)}
f25 = {c: v1 for c, d, v1, v0 in ler('DFC_MI_con', 2025)}; f24 = {c: v0 for c, d, v1, v0 in ler('DFC_MI_con', 2025)}
base = [('Ativo circulante', 2158236, 1569425, 'BP (CVM 1.01)'), ('Passivo circulante', p25['2.01'], p24['2.01'], 'CVM 2.01'),
        ('Estoques', 6139, 9766, 'BP'), ('Caixa e equivalentes', 229508, 295548, 'BP'), ('Aplicacoes financeiras CP', 1395510, 761939, 'BP'),
        ('Ativo total', 8668015, 7437571, 'BP'), ('Passivo nao circulante', p25['2.02'], p24['2.02'], 'CVM 2.02'),
        ('Patrimonio liquido', 4330716, 4235277, 'BP'), ('Debentures (CP + LP)', 3105016, 1792321, 'NE 9.1'),
        ('Receita liquida', 3157609, 3264554, 'DRE'), ('Lucro operacional', 721579, 948220, 'DRE'), ('D&A', 721077, 694816, 'NE 19'),
        ('Lucro liquido', 638352, 437498, 'DRE'), ('Juros sobre debentures', 281098, 84840, 'NE 20'),
        ('Caixa liquido das atividades operacionais (FCO)', f25['6.01'], f24['6.01'], 'DFC 6.01'),
        ('Adicoes ao imobilizado e intangivel (capex caixa)', -f25['6.02.02'], -f24['6.02.02'], 'DFC 6.02.02 (sinal invertido)'),
        ('Valor de mercado em 31/12 (R$ 11,32 x 292.977.928 acoes em circulacao)', '=11.32*292977.928', None, 'F01 RA p.28 (R$ 3,3 bi) e NE 15.1 p.86')]
for n, a, b, f in base:
    ws.append([n, a, b, f])
fmt(ws, 5, 21, 2, 3)
ws.cell(row=23, column=1, value='Indicadores').font = Font(bold=True)
ind = [('Liquidez corrente (AC / PC)', '=B5/B6', '=C5/C6', '0.00"x"'), ('Liquidez seca ((AC - estoques) / PC)', '=(B5-B7)/B6', '=(C5-C7)/C6', '0.00"x"'),
       ('Liquidez imediata ((caixa + aplic. CP) / PC)', '=(B8+B9)/B6', '=(C8+C9)/C6', '0.00"x"'),
       ('Endividamento geral ((PC + PNC) / AT)', '=(B6+B11)/B10', '=(C6+C11)/C10', PCT), ('Debentures / PL', '=B13/B12', '=C13/C12', '0.00"x"'),
       ('Divida liquida simples (debentures - caixa - aplic. CP)', '=B13-B8-B9', '=C13-C8-C9', NUM), ('EBITDA (lucro operacional + D&A)', '=B15+B16', '=C15+C16', NUM),
       ('Divida liquida simples / EBITDA', '=B29/B30', '=C29/C30', '0.00"x"'), ('Cobertura de juros (EBITDA / juros de debentures)', '=B30/B18', '=C30/C18', '0.0"x"'),
       ('Margem EBITDA', '=B30/B14', '=C30/C14', PCT), ('Margem liquida', '=B17/B14', '=C17/C14', PCT), ('ROE (lucro liquido / PL final)', '=B17/B12', '=C17/C12', PCT),
       ('FCO - capex (fluxo de caixa livre simplificado)', '=B19-B20', '=C19-C20', NUM),
       ('Valor de mercado / PL (< 1 = indicativo de impairment, CPC 01 item 12(d))', '=B21/B12', '', '0.00"x"')]
r = 24
for n, a, b, nf in ind:
    ws.cell(row=r, column=1, value=n); ws.cell(row=r, column=2, value=a).number_format = nf; ws.cell(row=r, column=3, value=b).number_format = nf
    for c in range(1, 4):
        ws.cell(row=r, column=c).border = BORDER
    r += 1

# ===================================================== 9. Grupo (PT 600)
ws = wb.create_sheet('Grupo_Componentes')
titulo(ws, 'Mapa do grupo e dos componentes (insumo PT 600)')
cab(ws, 4, ['Componente', 'Natureza / tratamento contabil', 'Ativo total', 'PL', 'Participacao', 'Fonte'], [40, 60, 14, 14, 12, 30])
for row in [('PetroReconcavo S.A. (controladora)', 'Entidade principal; concessionaria e operadora', 8597497, 4330716, '-', 'F01 BP p.51'),
            ('SPE Tieta Ltda. (ex-Maha Energy Brasil Ltda.)', 'Controlada integral: consolidacao; MEP na controladora; mais-valia amortizada R$ 37.602 em 2025', 879674, 803390, '100%', 'F01 NE 1.1 e 6 p.57/65'),
            ('Consorcios de E&P (14 dos 58 campos)', 'Operacoes em conjunto: reconhecimento proporcional', None, None, 'varia', 'F01 NE 1 p.57'),
            ('JOA midstream Brava (UPGN Guamare, RN)', 'Operacao em conjunto (CPC 19): 50%; aquisicao US$ 65 mi em 30/09/2025; operadora Brava', None, None, '50%', 'F01 NE 7 (i) p.70'),
            ('Consorcio/JOA Mandacaru (7 concessoes, RN)', 'Farm-out de 50% em 10/10/2025; operadora Mandacaru; ativos mantidos no imobilizado', None, None, '50%', 'F01 NE 7.4 p.71')]:
    ws.append(list(row))
fmt(ws, 5, 9, 3, 4); wrap(ws, 5, 9, 1, 2)
for i, t in enumerate(['Eliminacoes: investimento de 847.106 (controladora) eliminado; transacoes intragrupo 2025: R$ 31.667 (servicos/vendas) + R$ 44.763 (rateios) - NE 16.',
                       'Ativo da SPE Tieta = 10,1% do ativo consolidado; receita liquida consolidada menos a da controladora = R$ 444.831 (14,1%), antes de eliminacoes.',
                       'Auditor: a EY audita controladora e controlada (F12 9.1); nao ha auditor de componente distinto identificado nas fontes publicas.']):
    ws.cell(row=11 + i, column=1, value=t)

# ===================================================== 10. Achados preliminares
ws = wb.create_sheet('Achados_Preliminares')
titulo(ws, 'Achados preliminares do diagnostico - HIPOTESES a investigar (nao sao constatacoes de auditoria)')
cab(ws, 4, ['#', 'Observacao', 'Evidencia (fonte / pagina)', 'Norma / PT de destino', 'Natureza'], [4, 90, 38, 24, 18])
ach = [
 (1, 'Indicativo externo de impairment: valor de mercado (R$ 3,3 bi) inferior ao PL (R$ 4,33 bi) em 31/12/2025, com Brent medio 14% menor e lucro operacional 24% menor; a NE 3.4 afirma que nao houve indicativos. Avaliar o teste de recuperabilidade (CPC 01, item 12(d)).',
  'F01 RA p.28 (valor de mercado); NE 3.4 p.62; DRE p.52', 'NBC TA 540/315/330 (PT 540, 315)', 'Risco significativo'),
 (2, 'Divergencia na 4a emissao: a tabela NE 9.1 mostra serie 1 = 225.208 e serie 2 = 525.492, mas o texto da NE 9 e a NE 25 informam 1a serie R$ 525 mi e 2a serie R$ 225 mi; o texto cita 225.500 debentures da 2a serie (total 750.500 vs. R$ 750.000). Possivel troca de rotulos.',
  'F01 NE 9.1 p.72; NE 9 p.74-75; NE 25 p.101', 'NBC TA 450/500 (PT 450, 500)', 'Divulgacao'),
 (3, 'Data de liquidacao da 4a emissao: 18/12/2025 (NE 9 e NE 21) x 29/12/2025 (NE 25). Taxas: IPCA + 7,40% (NE 9) x 7,4564% e 7,2833% (NE 25).', 'F01 NE 9 p.74; NE 25 p.101', 'NBC TA 450/560', 'Divulgacao'),
 (4, 'A NE 25 "Evento subsequente" descreve fatos de dezembro/2025 (anteriores a data do balanco); apenas o swap de 12/01/2026 e evento subsequente (sem ajuste).', 'F01 NE 25 p.101; F23', 'NBC TA 560 (PT 560)', 'Divulgacao'),
 (5, 'Valores a pagar por aquisicoes - UPGN Guamare: R$ 18.515 descrito como US$ 3.559 (NE 11.1 e NE 21) e como US$ 9.952 (NE 11 a); com a PTAX de 31/12/2025 (5,5024) equivaleria a ~US$ 3.365.', 'F01 NE 11 p.79-80; NE 21 p.98', 'NBC TA 450/500', 'Divulgacao / mensuracao'),
 (6, 'As transferencias da movimentacao do imobilizado nao zeram (custo -81.390; depreciacao -62.640); a nota (a) atribui a creditos tributarios e itens de alto giro, sem quantificar.', 'F01 NE 7.1 p.69', 'NBC TA 500/520', 'Analitico'),
 (7, 'Troca de auditor (Deloitte -> EY) em 2025: auditoria inicial; o PAA de 2024 (receita) nao foi repetido em 2025, cujo PAA e a provisao para abandono.', 'F01 p.46-47; F07 p.37 e 42-45; F12 9.1', 'NBC TA 510/300/240', 'Contexto'),
 (8, 'Creditos de 2022 em disputa (gas) e arbitragem CCI confidencial com pedido contraposto: possivel passivo contingente e incerteza sobre o recebivel de R$ 55.917.', 'F01 NE 5.1 p.64; NE 12.3 p.82', 'NBC TA 501/540/250', 'Risco'),
 (9, 'Dividendos de R$ 300 mi declarados em 18/12/2025 com base na Lei 15.270/2025 (tributacao de dividendos a partir de 2026), R$ 200 mi no passivo nao circulante: verificar conformidade legal e as restricoes de covenants (distribuicao acima de 25% do lucro).', 'F01 NE 15.6 p.88; NE 9.3 p.75; BP', 'NBC TA 250/570', 'Leis e regulamentos'),
 (10, 'A NE 16.2 cita AGO de "24 de abril de 2024" fixando a remuneracao de 2025: provavel erro de data (AGO de 2025).', 'F01 NE 16.2 p.91', 'NBC TA 550/450', 'Divulgacao (menor)'),
 (11, 'Concentracao de receita: tres clientes = 83% da receita. Risco de credito, de corte e de dependencia; foi o PAA do auditor anterior.', 'F01 NE 21 p.97', 'NBC TA 240/315/505', 'Risco'),
]
for a in ach:
    ws.append(list(a))
wrap(ws, 5, 4 + len(ach), 1, 5)

os.makedirs('planilhas', exist_ok=True)
wb.save('planilhas/DIAG_RECV3.xlsx')
print('ok planilhas/DIAG_RECV3.xlsx | abas:', len(wb.sheetnames), '| comparativos com diferenca:', NDIF)
