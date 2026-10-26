# -*- coding: utf-8 -*-
# Gera o dossie em PDF (com sumario e referencias cruzadas clicaveis) a partir dos papeis de trabalho em Markdown.
# Preparador: matricula 211059370 | 24/10/2026
# Requisitos: pip install markdown playwright openpyxl ; python -m playwright install chromium (se o Chromium nao existir)
# Uso (na raiz do repositorio): python dossie/gerar_dossie.py
import re, glob, os, html, datetime
import markdown
from openpyxl import load_workbook

ORDEM = ['EC', '200', '210', '220', '230', '240', '250', '260', '265', '300', '315', '320', '330', '402', '450', '500',
         '501', '505', '510', '520', '530', '540', '550', '560', '570', '580', '600', '999']
SAIDA_HTML = 'dossie/DOSSIE_P2_RECV3_211059370.html'
SAIDA_PDF = 'dossie/DOSSIE_P2_RECV3_211059370.pdf'

def prep(txt):
    """Ajustes para o python-markdown: linha em branco antes de listas e recuo de 4 espacos em sublistas."""
    out = []; prev = ''
    for l in txt.splitlines():
        m = re.match(r'^( *)([-*]|\d+\.)\s', l)
        if m:
            ind = len(m.group(1))
            if ind and ind % 4: l = ' ' * (4 * ((ind + 1) // 2)) + l.lstrip()
            if prev.strip() and not re.match(r'^\s*([-*]|\d+\.)\s', prev) and not prev.startswith('>') and not prev.startswith('|'):
                out.append('')
        out.append(l); prev = l
    return '\n'.join(out)

def md2html(txt):
    return markdown.markdown(prep(txt), extensions=['tables', 'fenced_code', 'sane_lists'])

def linkar_pts(h, atual=None):
    """Transforma 'PT 315' (e 'PT EC') em links internos, sem mexer em atributos/links ja existentes."""
    def rep(m):
        cod = m.group(1)
        if cod == atual: return m.group(0)
        return f'<a class="xref" href="#pt-{cod.lower()}">PT {cod}</a>'
    partes = re.split(r'(<[^>]+>)', h); dentro_a = False; out = []
    for p in partes:
        if p.startswith('<a '): dentro_a = True
        elif p.startswith('</a'): dentro_a = False
        if p.startswith('<') or dentro_a: out.append(p); continue
        out.append(re.sub(r'\bPT[\s ](EC|\d{3})\b', rep, p))
    return ''.join(out)

# ---------- capa, identificacao ----------
secoes = []; sumario = []
capa = f'''<section class="capa">
<div class="topo">Universidade de Brasília · Auditoria 1 · 2026.2 · Prova 2</div>
<h1>Dossiê de Auditoria</h1>
<h2>PetroReconcavo S.A. (B3: RECV3)</h2>
<p class="sub">Demonstrações financeiras individuais e consolidadas – exercício findo em 31/12/2025</p>
<p class="sub">Cobertura: NBC TA Estrutura Conceitual e NBC TA 200 a 600</p>
<table class="ident"><tr><td>Preparador</td><td>Matrícula 211059370</td></tr>
<tr><td>Versão</td><td>v1.0 – 25/10/2026 (gerado em {datetime.date.today().strftime('%d/%m/%Y')})</td></tr>
<tr><td>Revisão</td><td>______ (campo reservado; não preenchido – não se simula revisão de terceiros)</td></tr></table>
<p class="aviso">Trabalho acadêmico. Usa apenas fontes públicas e dados <b>sintéticos</b> identificados como tais. Procedimentos classificados como
<b>[REALIZADO]</b> (evidência pública), <b>[PLANEJADO]</b> (falta de acesso) ou <b>[SIMULADO]</b> (fins didáticos). <b>Não contém opinião de auditoria.</b>
Nenhuma empresa, auditor, cliente ou fornecedor foi contatado.</p></section>'''
secoes.append(capa)

def bloco(id_, titulo, corpo_html, nivel_sum=1):
    sumario.append((id_, titulo, nivel_sum))
    secoes.append(f'<section class="pt" id="{id_}"><div class="faixa">{html.escape(titulo)}</div>{corpo_html}</section>')

# ---------- parte I: guia ----------
guia = '''<h1>Guia de leitura</h1>
<p>O dossiê segue o índice do <a class="xref" href="#pt-230">PT 230</a>. Cada papel de trabalho traz: código, título, norma e itens; objetivo;
afirmações e riscos; população, fontes e critérios; procedimentos; evidências e referências cruzadas; resultados, exceções, conclusão e limitações;
preparador (matrícula e data) e campo de revisão em branco.</p>
<table><thead><tr><th>Código</th><th>Significado</th><th>Onde</th></tr></thead><tbody>
<tr><td>F01…F39</td><td>Fontes (URL, data e hora de acesso, página do PDF)</td><td>Registro de fontes (anexo A)</td></tr>
<tr><td>R01…R17</td><td>Riscos por afirmação</td><td><a class="xref" href="#pt-315">PT 315</a></td></tr>
<tr><td>P01…P17</td><td>Respostas</td><td><a class="xref" href="#pt-330">PT 330</a></td></tr>
<tr><td>D01…D07</td><td>Deficiências de controle</td><td><a class="xref" href="#pt-265">PT 265</a></td></tr>
<tr><td>M01…M12; D-I1…D-I5</td><td>Distorções de valor e de divulgação</td><td><a class="xref" href="#pt-450">PT 450</a></td></tr>
<tr><td>C01…C09</td><td>Confirmações externas (modelos)</td><td><a class="xref" href="#pt-505">PT 505</a></td></tr>
<tr><td>T01…T10; G01…G06</td><td>Testes de 100% na base do imobilizado e da folha</td><td><a class="xref" href="#pt-530">PT 530</a>; <a class="xref" href="#pt-450">PT 450</a></td></tr>
</tbody></table>
<p>Arquivos citados (planilhas, scripts, base) estão no repositório; ver o README.</p>'''
bloco('guia', 'Guia de leitura', guia)

# ---------- matriz de cobertura (resumo) e versoes ----------
try:
    wb = load_workbook('matriz_cobertura/MATRIZ_COBERTURA_NORMAS.xlsx')
    ws = wb['Matriz_por_norma']; linhas = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0] or not str(r[0]).strip() or str(r[0]).startswith('Cobertura') or str(r[0]).startswith('Outros'): continue
        cob = r[8]; cob = f'{cob:.0%}' if isinstance(cob, float) else cob
        linhas.append(f'<tr><td>{r[0]}</td><td>{html.escape(str(r[1]))}</td><td>{html.escape(str(r[2]))}</td>'
                      f'<td><a class="xref" href="#pt-{str(r[0]).lower()}">{r[3]}</a></td><td>{r[6]}</td><td>{r[7]}</td><td>{cob}</td><td>{r[12]}</td></tr>')
    mat = ('<h1>Matriz de cobertura das normas (resumo)</h1><p>Itens de requisito numerados de cada norma (extraídos do consolidado da disciplina) '
           'confrontados com o PT correspondente. O detalhe item a item e as análises de “não se aplica” estão na planilha '
           '<code>matriz_cobertura/MATRIZ_COBERTURA_NORMAS.xlsx</code> (abas <i>Itens_requisitos</i> e <i>Nao_se_aplica</i>).</p>'
           '<table><thead><tr><th>Norma</th><th>Título</th><th>Versão</th><th>PT</th><th>Itens de requisito</th><th>Citados</th><th>Cobertura</th><th>“Não se aplica” analisados</th></tr></thead><tbody>'
           + ''.join(linhas) + '</tbody></table>')
    wv = wb['Versoes_normas']; lv = []
    for r in wv.iter_rows(min_row=5, values_only=True):
        if r[0] and r[1]: lv.append(f'<tr><td>{r[0]}</td><td>{html.escape(str(r[1]))}</td><td>{r[2]}</td><td>{html.escape(str(r[3]))}</td></tr>')
    mat += ('<h2>Registro das versões das normas consultadas</h2><table><thead><tr><th>Norma</th><th>Sigla/versão</th><th>Data</th><th>Vigência / observação</th></tr></thead><tbody>'
            + ''.join(lv) + '</tbody></table><p class="nota">2009* = versão original de 2009; o texto do consolidado não traz a data de aprovação.</p>')
    bloco('matriz', 'Matriz de cobertura e versões das normas', mat)
except FileNotFoundError:
    print('aviso: matriz nao encontrada - rode codigo/gerar_matriz_cobertura.py')

# ---------- PTs ----------
for cod in ORDEM:
    fs = glob.glob(f'papeis_trabalho/PT_{cod}_*.md')
    if not fs: print('aviso: falta PT', cod); continue
    txt = open(fs[0], encoding='utf-8').read()
    titulo = re.search(r'^#\s+(.+)$', txt, re.M).group(1).strip()
    h = linkar_pts(md2html(txt), atual=cod)
    h = h.replace('<td>[REALIZADO', '<td><span class="tag r">[REALIZADO</span>')
    for k, c in (('REALIZADO', 'r'), ('SIMULADO', 's'), ('PLANEJADO', 'p')):
        h = re.sub(r'\[(%s[^\]]*)\]' % k, r'<span class="tag %s">[\1]</span>' % c, h)
    bloco(f'pt-{cod.lower()}', titulo, h)

# ---------- anexos ----------
try:
    wb = load_workbook('planilhas/REG_FONTES_RECV3.xlsx', read_only=True); ws = wb['Registro_Fontes']; lin = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        url = str(r[5] or '')
        lin.append(f'<tr><td>{r[0]}</td><td>{html.escape(str(r[2] or ""))}</td><td>{html.escape(str(r[3] or ""))}</td>'
                   f'<td class="url">{html.escape(url)}</td><td>{r[6] or ""} {r[7] or ""}</td><td>{html.escape(str(r[10] or ""))[:400]}</td></tr>')
    ane = ('<h1>Anexo A – Relação de fontes</h1><p>Extrato do <code>planilhas/REG_FONTES_RECV3.xlsx</code> (a planilha completa traz também a base do horário, '
           'o arquivo local, o uso previsto, o SHA-256 e o tamanho). Horário de Brasília (UTC−3).</p><table class="fontes"><thead><tr><th>ID</th><th>Documento</th>'
           '<th>Exercício</th><th>URL</th><th>Acesso (data e hora)</th><th>Páginas / seções</th></tr></thead><tbody>' + ''.join(lin) + '</tbody></table>')
    bloco('anexo-a', 'Anexo A – Relação de fontes', ane)
except FileNotFoundError:
    print('aviso: REG_FONTES nao encontrado')
for id_, tit, arq in [('anexo-b', 'Anexo B – Base simulada (README)', 'base_simulada/README_BASE_SIMULADA.md'),
                      ('anexo-c', 'Anexo C – Diagnóstico da companhia (etapa 2)', 'diagnostico/ETAPA2_Diagnostico_RECV3.md'),
                      ('anexo-d', 'Anexo D – Registro do uso de IA', 'docs/USO_DE_IA.md')]:
    if os.path.exists(arq):
        bloco(id_, tit, linkar_pts(md2html(open(arq, encoding='utf-8').read())))

# ---------- sumario ----------
PAGINAS = {}
if os.path.exists('dossie/_paginas.txt'):
    for l in open('dossie/_paginas.txt', encoding='utf-8'):
        k, v = l.rstrip('\n').split('\t'); PAGINAS[k] = v
som = '<section class="sumario"><h1>Sumário</h1><table class="som">' + ''.join(
    f'<tr><td><a href="#{i}">{html.escape(t)}</a></td><td class="pg">{PAGINAS.get(i, "")}</td></tr>' for i, t, _ in sumario) + '</table></section>'
secoes.insert(1, som)

CSS = '''
@page { size: A4; margin: 18mm 14mm 18mm 14mm; }
body { font-family: "DejaVu Sans", Arial, sans-serif; font-size: 9.2pt; color: #1b1b1b; line-height: 1.38; }
section { page-break-before: always; } section.capa { page-break-before: auto; }
.capa { text-align: center; padding-top: 50mm; } .capa h1 { font-size: 30pt; color: #1F4E78; margin: 6mm 0 2mm; }
.capa h2 { font-size: 18pt; margin: 0 0 6mm; } .capa .sub { font-size: 11pt; margin: 1mm; } .capa .topo { font-size: 10pt; color: #555; }
.capa table.ident { margin: 14mm auto; border-collapse: collapse; font-size: 10pt; } .capa table.ident td { border: 1px solid #bbb; padding: 4px 10px; text-align: left; }
.capa .aviso { margin: 10mm 18mm; font-size: 9pt; border: 1px solid #1F4E78; padding: 8px; text-align: justify; }
.faixa { font-size: 8pt; color: #fff; background: #1F4E78; padding: 3px 6px; margin-bottom: 4mm; }
h1 { font-size: 15pt; color: #1F4E78; border-bottom: 2px solid #1F4E78; padding-bottom: 2px; }
h2 { font-size: 12pt; color: #1F4E78; margin-top: 5mm; } h3 { font-size: 10.5pt; }
table { border-collapse: collapse; width: 100%; margin: 2mm 0 3mm; font-size: 8pt; page-break-inside: auto; }
tr { page-break-inside: avoid; } th { background: #DDEBF7; } th, td { border: 1px solid #BFBFBF; padding: 2.5px 4px; vertical-align: top; text-align: left; }
blockquote { border-left: 3px solid #1F4E78; margin: 2mm 0; padding: 1mm 4mm; background: #F5F8FC; font-size: 8.6pt; }
code { font-size: 8pt; background: #f3f3f3; padding: 0 2px; } pre { white-space: pre-wrap; font-size: 7.5pt; background: #f6f6f6; padding: 3mm; }
a { color: #1F4E78; text-decoration: none; } a.xref { border-bottom: 1px dotted #1F4E78; }
.tag { font-weight: bold; font-size: 7.6pt; padding: 0 2px; border-radius: 2px; } .tag.r { color: #1E6B34; } .tag.s { color: #8A5A00; } .tag.p { color: #8B1C1C; }
.sumario table.som td { border: none; border-bottom: 1px dotted #ccc; padding: 1.5px 4px; font-size: 8.6pt; } .sumario td.pg { text-align: right; width: 12mm; } .nota { font-size: 8pt; color: #555; }
table.fontes td.url { word-break: break-all; font-size: 6.8pt; width: 28%; }
'''
doc = f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Dossiê P2 – RECV3 – 211059370</title><style>{CSS}</style></head><body>{"".join(secoes)}</body></html>'
os.makedirs('dossie', exist_ok=True)
open(SAIDA_HTML, 'w', encoding='utf-8').write(doc)

from playwright.sync_api import sync_playwright
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page()
    pg.goto('file://' + os.path.abspath(SAIDA_HTML)); pg.wait_for_timeout(500)
    pg.pdf(path=SAIDA_PDF, format='A4', print_background=True, display_header_footer=True,
           header_template='<div style="font-size:7px;width:100%;text-align:right;padding-right:14mm;color:#666">Dossiê P2 · PetroReconcavo S.A. (RECV3) · 2025 · matrícula 211059370</div>',
           footer_template='<div style="font-size:7px;width:100%;text-align:center;color:#666">Página <span class="pageNumber"></span> de <span class="totalPages"></span></div>',
           margin={'top': '16mm', 'bottom': '16mm', 'left': '12mm', 'right': '12mm'})
    b.close()
# segunda passagem: localiza a pagina de cada secao (texto da faixa) e regera com o numero no sumario
import subprocess
if not os.environ.get('DOSSIE_PASSO2'):
    pags = {}
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', SAIDA_PDF], capture_output=True, text=True).stdout).group(1))
    textos = [subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), SAIDA_PDF, '-'], capture_output=True, text=True).stdout for i in range(1, n + 1)]
    for i, t, _ in sumario:
        alvo = re.sub(r'\s+', ' ', t)[:40]
        for pg in range(2, n):
            if alvo in re.sub(r'\s+', ' ', textos[pg]): pags[i] = str(pg + 1); break
    with open('dossie/_paginas.txt', 'w', encoding='utf-8') as fh:
        for k, v in pags.items(): fh.write(f'{k}\t{v}\n')
    os.environ['DOSSIE_PASSO2'] = '1'
    subprocess.run(['python3', __file__], check=True)
else:
    print('ok', SAIDA_PDF)
